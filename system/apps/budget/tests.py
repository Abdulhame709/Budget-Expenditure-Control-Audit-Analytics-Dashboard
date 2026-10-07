"""PHASE 5 — Budget Management: calculations, constraints, workflow, perms.

Proves: annual = Σ(months) (clean + DB CHECK), non-negative, unique line
per (version, dept, account), expense accounts only, category inheritance,
one budget per FY, draft⇄approved workflow with audit, approved-version
lock (revision flow), DB-computed totals & the analysis-version selector,
budget.view/budget.edit gating over real HTTP, and the Budget-vs-Actual /
Variance / Out-of-Budget FOUNDATION (monthly_budget_map, is_budgeted).
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.test import Client, TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook

from apps.accounts.models import Role, UserRole
from apps.governance.models import AuditLog
from apps.budget import services, spreadsheet
from apps.budget.models import (
    MONTH_FIELDS,
    Budget,
    BudgetLine,
    BudgetTemplate,
    BudgetTemplateCell,
    BudgetTemplateColumn,
    BudgetTemplateRow,
    BudgetVersion,
)
from apps.budget.template_import import import_budget_workbook, refresh_budget_workbook
from apps.budget.template_runtime import SheetRuntime, template_monthly_output
from apps.reference.models import (
    Account,
    Currency,
    Department,
    ExpenseCategory,
    FiscalYear,
)

User = get_user_model()
PW = "S3cure-Demo-Pass!"
D = Decimal


def make_user(username, role_code):
    user = User.objects.create_user(username=username, password=PW)
    UserRole.objects.create(user=user, role=Role.objects.get(code=role_code))
    return user


def make_fy(code="FY2026", year=2026):
    return FiscalYear.objects.create(
        code=code, name=f"السنة {year}", year=year,
        start_date=date(year, 1, 1), end_date=date(year, 12, 31),
    )


def make_budget(fy, name="ميزانية تجريبية"):
    return Budget.objects.create(fiscal_year=fy, name=name)


def make_version(budget, version=1, status="draft"):
    return BudgetVersion.objects.create(
        budget=budget, version=version, status=status)


def base_objects():
    dept = Department.objects.create(code="OPS1", name="العمليات")
    cat = ExpenseCategory.objects.create(code="TRV", name="تنقّل")
    acc = Account.objects.create(code="5101", name="مصروفات تنقّل",
                                 account_type="expense", expense_category=cat)
    asset = Account.objects.create(code="1100", name="نقدية",
                                   account_type="asset")
    return dept, cat, acc, asset


def make_line(version, dept, acc, months=None, **extra):
    months = months or {}
    values = {f: months.get(f, 0) for f in MONTH_FIELDS}
    line = BudgetLine(version=version, department=dept, account=acc,
                      **values, **extra)
    line.full_clean()   # derives annual_amount (as the views do) before save
    line.save()
    return line


def sample_workbook_bytes():
    workbook = Workbook()
    detail = workbook.active
    detail.title = "تفصيلي"
    detail.merge_cells("A1:D1")
    detail["A1"] = "نموذج الموازنة التفصيلي"
    detail["A2"] = "رقم الحساب"
    detail["B2"] = "اسم الحساب"
    detail["C2"] = "يناير"
    detail["D2"] = "الإجمالي السنوي"
    detail["A3"] = "5101"
    detail["B3"] = "مصروفات تنقّل"
    detail["C3"] = 100
    detail["D3"] = "=SUM(C3:C3)"
    summary = workbook.create_sheet("الإجماليات")
    summary["A1"] = "إجمالي الحسابات الرئيسية"
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    content.seek(0)
    return content.getvalue()


class BudgetTemplateImportTests(TestCase):
    def setUp(self):
        Currency.objects.update_or_create(
            code="YER",
            defaults={
                "name_ar": "ريال يمني", "name_en": "Yemeni Rial",
                "exchange_rate_to_base": 1, "is_base": True, "is_active": True,
            },
        )
        self.account = Account.objects.create(
            code="5101", name="مصروفات تنقّل", account_type="expense",
        )
        self.analytical_account = Account.objects.create(
            code="510101", name="تنقّل تحليلي", account_type="expense",
            level=6, ledger_type="sub", parent=self.account,
        )
        self.user = User.objects.create_superuser(
            username="template-admin", email="template@example.com", password=PW,
        )

    def test_import_preserves_sheets_formula_and_account_mapping(self):
        upload = SimpleUploadedFile(
            "budget.xlsx", sample_workbook_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        template = import_budget_workbook(upload, name="نموذج اختباري", user=self.user)
        self.assertEqual(template.currency.code, "YER")
        self.assertEqual(template.sheets.count(), 2)
        detail = template.sheets.get(name="تفصيلي")
        self.assertEqual(detail.merged_ranges, ["A1:D1"])
        self.assertEqual(detail.rows.get(row_number=3).account, self.account)
        self.assertEqual(
            detail.rows.get(row_number=3).cells.get(coordinate="D3").formula,
            "=SUM(C3:C3)",
        )
        self.assertEqual(
            template.sheets.get(name="الإجماليات").purpose,
            "summary",
        )

    def test_runtime_recalculates_formula_after_edit_and_builds_monthly_output(self):
        template = import_budget_workbook(
            SimpleUploadedFile("budget.xlsx", sample_workbook_bytes()),
            name="نموذج حي", user=self.user,
        )
        detail = template.sheets.get(name="تفصيلي")
        row = detail.rows.get(row_number=3)
        row.main_account = self.account
        row.analytical_account = self.analytical_account
        row.save()
        january = detail.columns.get(column_letter="C")
        january.role = "month_01"
        january.save(update_fields=["role"])
        annual = detail.columns.get(column_letter="D")
        annual.role = BudgetTemplateColumn.ROLE_ANNUAL
        annual.save(update_fields=["role"])
        input_cell = row.cells.get(coordinate="C3")
        formula_cell = row.cells.get(coordinate="D3")
        self.assertEqual(SheetRuntime(detail).value("D3"), Decimal("100"))
        input_cell.override_value = "250"
        input_cell.save(update_fields=["override_value"])
        self.assertEqual(SheetRuntime(detail).value("D3"), Decimal("250"))
        output = template_monthly_output(detail)
        self.assertEqual(len(output), 1)
        self.assertEqual(output[0]["months"][1], Decimal("250.00"))
        self.assertEqual(output[0]["annual_amount"], Decimal("250.00"))

    def test_refresh_preserves_account_mapping_and_replaces_source_values(self):
        template = import_budget_workbook(
            SimpleUploadedFile("budget.xlsx", sample_workbook_bytes()),
            name="نموذج متجدد", user=self.user,
        )
        row = template.sheets.get(name="تفصيلي").rows.get(row_number=3)
        row.main_account = self.account
        row.analytical_account = self.analytical_account
        row.save()
        updated = Workbook()
        sheet = updated.active
        sheet.title = "تفصيلي"
        sheet["A1"] = "نموذج الموازنة التفصيلي"
        sheet["A2"] = "رقم الحساب"
        sheet["B2"] = "اسم الحساب"
        sheet["C2"] = "يناير"
        sheet["D2"] = "الإجمالي السنوي"
        sheet["A3"] = "5101"
        sheet["B3"] = "مصروفات تنقّل"
        sheet["C3"] = 400
        sheet["D3"] = "=SUM(C3:C3)"
        content = BytesIO()
        updated.save(content)
        updated.close()
        refresh_budget_workbook(
            template,
            SimpleUploadedFile("updated.xlsx", content.getvalue()),
            user=self.user,
        )
        row.refresh_from_db()
        self.assertEqual(row.main_account, self.account)
        self.assertEqual(row.analytical_account, self.analytical_account)
        self.assertEqual(row.cells.get(coordinate="C3").raw_value, "400")

    def test_total_row_can_sum_previous_data_rows(self):
        template = import_budget_workbook(
            SimpleUploadedFile("budget.xlsx", sample_workbook_bytes()),
            name="نموذج إجمالي مرن", user=self.user,
        )
        detail = template.sheets.get(name="تفصيلي")
        annual = detail.columns.get(column_letter="D")
        annual.role = BudgetTemplateColumn.ROLE_ANNUAL
        annual.save(update_fields=["role"])
        source = detail.rows.get(row_number=3)
        second = BudgetTemplateRow.objects.create(
            sheet=detail, row_number=4, position=4,
            row_type=BudgetTemplateRow.TYPE_DATA, label="بيانات ثانية",
        )
        BudgetTemplateCell.objects.create(
            row=second, column=annual, coordinate="D4", raw_value="50",
            is_editable=True,
        )
        total = BudgetTemplateRow.objects.create(
            sheet=detail, row_number=5, position=5,
            row_type=BudgetTemplateRow.TYPE_SUBTOTAL, label="إجمالي فرعي",
            is_included=False,
        )
        BudgetTemplateCell.objects.create(
            row=total, column=annual, coordinate="D5", is_editable=True,
        )
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("budget:template_sheet", args=[detail.pk]),
            {"action": "sum_previous", "row_id": total.pk},
        )
        self.assertEqual(response.status_code, 302)
        total_cell = total.cells.get(column=annual)
        self.assertEqual(total_cell.formula, "=SUM(D3,D4)")
        self.assertEqual(SheetRuntime(detail).value("D5"), Decimal("150"))
        self.assertEqual(source.position, 3)

    def test_template_rows_move_and_resequence_after_delete(self):
        template = import_budget_workbook(
            SimpleUploadedFile("budget.xlsx", sample_workbook_bytes()),
            name="نموذج ترتيب", user=self.user,
        )
        detail = template.sheets.get(name="تفصيلي")
        rows = list(detail.rows.order_by("position"))
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("budget:template_sheet", args=[detail.pk]),
            {"action": "move_up", "row_id": rows[-1].pk},
        )
        self.assertEqual(response.status_code, 302)
        ordered = list(detail.rows.order_by("position").values_list("pk", flat=True))
        self.assertEqual(ordered[-2:], [rows[-1].pk, rows[-2].pk])
        response = self.client.post(
            reverse("budget:template_sheet", args=[detail.pk]),
            {"action": "delete_row", "row_id": rows[0].pk},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            list(detail.rows.order_by("position").values_list("position", flat=True)),
            list(range(1, detail.rows.count() + 1)),
        )

    def test_template_pages_and_upload_require_expected_permissions(self):
        client = Client()
        client.force_login(self.user)
        response = client.post(
            reverse("budget:template_import"),
            {
                "name": "مرفوع من الواجهة",
                "description": "اختبار",
                "workbook": SimpleUploadedFile(
                    "ui-budget.xlsx", sample_workbook_bytes(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )
        template = BudgetTemplate.objects.get(name="مرفوع من الواجهة")
        self.assertRedirects(
            response, reverse("budget:template_detail", args=[template.pk])
        )
        self.assertEqual(client.get(reverse("budget:template_list")).status_code, 200)
        self.assertEqual(
            client.get(reverse("budget:template_sheet", args=[template.sheets.first().pk])).status_code,
            200,
        )

    def _sales_cost_sheet(self):
        template = import_budget_workbook(
            SimpleUploadedFile("budget.xlsx", sample_workbook_bytes()),
            name="نموذج محرر الجداول", user=self.user,
        )
        detail = template.sheets.get(name="تفصيلي")
        detail.name = "تفصيلي تكلفة مبيعات"
        detail.save(update_fields=["name"])
        row = detail.rows.get(row_number=3)
        row.main_account = self.account
        row.analytical_account = self.analytical_account
        row.save(update_fields=["main_account", "analytical_account"])
        return detail, row

    def test_spreadsheet_page_and_snapshot_preserve_layout(self):
        detail, _ = self._sales_cost_sheet()
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("budget:template_sheet_spreadsheet", args=[detail.pk])
        )
        self.assertEqual(response.status_code, 200)
        response = self.client.get(
            reverse("budget:template_sheet_spreadsheet_data", args=[detail.pk])
        )
        self.assertEqual(response.status_code, 200)
        workbook = response.json()["workbook"]
        worksheet = next(iter(workbook["sheets"].values()))
        self.assertEqual(worksheet["mergeData"][0]["endColumn"], 3)
        self.assertEqual(worksheet["cellData"]["2"]["3"]["f"], "=SUM(C3:C3)")
        self.assertEqual(worksheet["columnData"]["0"]["custom"]["budgetColumnId"], detail.columns.get(column_index=1).pk)

    def test_spreadsheet_save_syncs_structure_and_keeps_account_links(self):
        detail, linked_row = self._sales_cost_sheet()
        workbook = spreadsheet.workbook_snapshot(detail)
        worksheet = next(iter(workbook["sheets"].values()))
        linked_metadata = worksheet["rowData"][2]
        linked_cells = worksheet["cellData"][2]
        linked_cells[2]["v"] = 250
        linked_cells[3]["f"] = "=SUM(C2:C2)"
        worksheet["rowData"] = {
            0: worksheet["rowData"][0],
            1: linked_metadata,
            2: {"h": 31},
        }
        worksheet["cellData"] = {
            0: worksheet["cellData"][0],
            1: linked_cells,
            2: {0: {"v": "بند جديد"}, 2: {"v": 75}},
        }
        worksheet["columnData"][0]["w"] = 190
        worksheet["mergeData"] = [{
            "startRow": 0, "endRow": 0, "startColumn": 0, "endColumn": 1,
        }]

        result = spreadsheet.save_workbook_snapshot(detail, workbook)

        self.assertEqual(result["rows"], 3)
        self.assertEqual(detail.rows.count(), 3)
        linked_row.refresh_from_db()
        self.assertEqual(linked_row.position, 2)
        self.assertEqual(linked_row.main_account, self.account)
        self.assertEqual(linked_row.analytical_account, self.analytical_account)
        changed = linked_row.cells.get(column__column_index=3)
        self.assertEqual(changed.effective_value, "250")
        self.assertEqual(detail.columns.get(column_index=1).width, Decimal("190"))
        detail.refresh_from_db()
        self.assertEqual(detail.merged_ranges, ["A1:B1"])
        self.assertEqual(detail.rows.get(position=3).height, Decimal("31"))

    def test_spreadsheet_save_requires_edit_permission(self):
        detail, _ = self._sales_cost_sheet()
        viewer = make_user("spreadsheet-viewer", "viewer")
        client = Client()
        client.force_login(viewer)
        response = client.get(
            reverse("budget:template_sheet_spreadsheet", args=[detail.pk])
        )
        self.assertEqual(response.status_code, 200)
        response = client.post(
            reverse("budget:template_sheet_spreadsheet_save", args=[detail.pk]),
            data={"workbook": spreadsheet.workbook_snapshot(detail)},
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)

    def test_spreadsheet_trial_is_limited_to_sales_cost_sheet(self):
        template = import_budget_workbook(
            SimpleUploadedFile("budget.xlsx", sample_workbook_bytes()),
            name="نموذج غير تجريبي", user=self.user,
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse(
            "budget:template_sheet_spreadsheet", args=[template.sheets.get(name="تفصيلي").pk],
        ))
        self.assertEqual(response.status_code, 404)


# ================================================================ calculations
class BudgetLineCalculationTests(TestCase):
    def setUp(self):
        self.fy = make_fy()
        self.budget = make_budget(self.fy)
        self.version = make_version(self.budget)
        self.dept, self.cat, self.acc, self.asset = base_objects()

    def test_annual_is_derived_from_months(self):
        line = BudgetLine(
            version=self.version, department=self.dept, account=self.acc,
            m01=D("1000.50"), m02=D("2000.25"), m06=D("500.00"),
        )
        line.full_clean()
        self.assertEqual(line.annual_amount, D("3500.75"))
        self.assertEqual(line.expense_category_id, self.cat.pk)  # inherited

    def test_db_check_rejects_mismatched_annual(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                BudgetLine.objects.create(
                    version=self.version, department=self.dept, account=self.acc,
                    m01=D("100"), annual_amount=D("999"),
                )

    def test_db_check_allows_correct_sum(self):
        line = make_line(self.version, self.dept, self.acc,
                         {"m01": D("100"), "m03": D("50")})
        self.assertEqual(line.annual_amount, D("150"))

    def test_negative_month_rejected(self):
        line = BudgetLine(
            version=self.version, department=self.dept, account=self.acc,
            m01=D("-1"),
        )
        with self.assertRaises(ValidationError):
            line.full_clean()

    def test_non_expense_account_rejected(self):
        line = BudgetLine(
            version=self.version, department=self.dept, account=self.asset,
            m01=D("10"),
        )
        with self.assertRaises(ValidationError) as ctx:
            line.full_clean()
        self.assertIn("account", ctx.exception.error_dict)

    def test_unique_line_per_version_dept_account(self):
        make_line(self.version, self.dept, self.acc, {"m01": D("10")})
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                # bypass full_clean to prove the DB-level constraint
                BudgetLine.objects.create(
                    version=self.version, department=self.dept, account=self.acc,
                    m01=D("20"), annual_amount=D("20"),
                )

    def test_one_budget_per_fiscal_year(self):
        with self.assertRaises(ValidationError) as ctx:
            Budget(fiscal_year=self.fy, name="أخرى").full_clean()
        self.assertIn("fiscal_year", ctx.exception.error_dict)


# ================================================================ workflow
class BudgetWorkflowTests(TestCase):
    def setUp(self):
        self.fy = make_fy()
        self.budget = make_budget(self.fy)
        self.version = make_version(self.budget)
        self.dept, self.cat, self.acc, _ = base_objects()
        make_line(self.version, self.dept, self.acc, {"m01": D("1000")})
        self.admin = make_user("adm_b5", "admin")

    def test_analysis_version_none_until_approved(self):
        self.assertIsNone(services.get_analysis_version(self.budget))
        self.version.status = "approved"
        self.version.save()
        analysis = services.get_analysis_version(self.budget)
        self.assertEqual(analysis.pk, self.version.pk)

    def test_latest_approved_wins_for_analysis(self):
        v2 = make_version(self.budget, version=2, status="approved")
        analysis = services.get_analysis_version(self.budget)
        self.assertEqual(analysis.pk, v2.pk)  # revised budget becomes THE one
        self.assertTrue(BudgetVersion.objects.get(pk=v2.pk).is_analysis_version)
        self.assertFalse(
            BudgetVersion.objects.get(pk=self.version.pk).is_analysis_version)


class VersionTotalsServiceTests(TestCase):
    """The numbers analysis will use — straight from DB aggregates."""

    def setUp(self):
        self.fy = make_fy()
        self.budget = make_budget(self.fy)
        self.v = make_version(self.budget)
        self.dept, self.cat, self.acc, _ = base_objects()
        self.dept2 = Department.objects.create(code="FIN1", name="المالية")
        self.acc2 = Account.objects.create(code="5102", name="مستلزمات",
                                           account_type="expense")
        make_line(self.v, self.dept, self.acc,
                  {"m01": D("100"), "m02": D("200")})
        make_line(self.v, self.dept, self.acc2, {"m03": D("50")})
        make_line(self.v, self.dept2, self.acc, {"m01": D("400")})

    def test_version_totals(self):
        totals = services.version_totals(self.v)
        self.assertEqual(totals["annual_total"], D("750"))
        self.assertEqual(totals["monthly"][0], D("500"))   # 100 + 400
        self.assertEqual(totals["monthly"][2], D("50"))
        self.assertEqual(totals["monthly"][3], D("0"))
        self.assertEqual(totals["line_count"], 3)

    def test_totals_by_department_and_account(self):
        by_dept = services.totals_by(self.v, "department")
        self.assertEqual(
            {r["key"]: r["total"] for r in by_dept},
            {self.dept.pk: D("350"), self.dept2.pk: D("400")},
        )
        by_acc = services.totals_by(self.v, "account")
        self.assertEqual(
            {r["key"]: r["total"] for r in by_acc},
            {self.acc.pk: D("700"), self.acc2.pk: D("50")},
        )

    def test_monthly_budget_map_foundation(self):
        m = services.monthly_budget_map(self.v)
        key = (self.dept.pk, self.acc.pk)
        self.assertIn(key, m)
        self.assertEqual(m[key]["monthly"][0], D("100"))
        self.assertEqual(m[key]["annual"], D("300"))

    def test_is_budgeted_foundation_oob(self):
        # covered pair
        self.assertTrue(services.is_budgeted(self.v, self.dept, self.acc))
        # another dept using the same account → NOT budgeted (OOB candidate)
        dept3 = Department.objects.create(code="X9", name="غير ممولة")
        self.assertFalse(services.is_budgeted(self.v, dept3, self.acc))
        self.assertIsNone(services.line_for(self.v, dept3, self.acc))


# ================================================================ HTTP + perms
class BudgetHttpTests(TestCase):
    def setUp(self):
        self.fy = make_fy()
        self.budget = make_budget(self.fy)
        self.version = make_version(self.budget)
        self.dept, self.cat, self.acc, _ = base_objects()
        make_line(self.version, self.dept, self.acc, {"m01": D("1000")})
        self.admin = make_user("adm_http5", "admin")
        self.finance = make_user("fin_http5", "finance")
        self.auditor = make_user("aud_http5", "auditor")
        self.mgmt = make_user("mgmt_http5", "management")

    def _client(self, user):
        c = Client()
        c.force_login(user)
        return c

    def test_anonymous_redirected(self):
        response = self.client.get(reverse("budget:budget_list"))
        self.assertEqual(response.status_code, 302)
        self.assertIn("next=", response.url)

    def test_management_denied_budget_pages(self):
        # management has no budget.view in the approved Phase-3 matrix
        c = self._client(self.mgmt)
        self.assertEqual(c.get(reverse("budget:budget_list")).status_code, 403)
        self.assertEqual(
            c.get(reverse("budget:version_detail", args=[self.version.pk])).status_code,
            403)

    def test_auditor_views_but_cannot_mutate(self):
        c = self._client(self.auditor)
        self.assertEqual(c.get(reverse("budget:budget_list")).status_code, 200)
        self.assertEqual(
            c.get(reverse("budget:version_detail", args=[self.version.pk])).status_code,
            200)
        self.assertEqual(
            c.get(reverse("budget:line_create", args=[self.version.pk])).status_code,
            403)
        self.assertEqual(
            c.post(reverse("budget:version_approve", args=[self.version.pk])).status_code,
            403)

    def test_finance_creates_budget_with_version1(self):
        c = self._client(self.finance)
        fy27 = make_fy(code="FY2027", year=2027)
        response = c.post(reverse("budget:budget_create"), {
            "fiscal_year": fy27.pk, "name": "ميزانية المالية", "notes": "",
        })
        self.assertEqual(response.status_code, 302)
        budget = Budget.objects.get(name="ميزانية المالية")
        v1 = budget.versions.get()
        self.assertEqual(v1.version, 1)
        self.assertEqual(v1.status, "draft")
        self.assertTrue(AuditLog.objects.filter(
            action="entity_created", entity_type="budget",
            entity_id=str(budget.pk)).exists())
        self.assertTrue(AuditLog.objects.filter(
            action="budget_version_created",
            entity_id=str(v1.pk)).exists())

    def test_finance_adds_line_total_derived(self):
        c = self._client(self.finance)
        dept2 = Department.objects.create(code="FIN1", name="المالية")
        response = c.post(
            reverse("budget:line_create", args=[self.version.pk]),
            {"department": dept2.pk, "account": self.acc.pk,
             "expense_category": "", "notes": "",
             "m01": "1500", "m06": "500"},
        )
        self.assertEqual(response.status_code, 302)
        line = BudgetLine.objects.get(version=self.version,
                                      department=dept2, account=self.acc)
        self.assertEqual(line.annual_amount, D("2000"))
        self.assertTrue(AuditLog.objects.filter(
            action="entity_created", entity_type="budgetline",
            entity_id=str(line.pk)).exists())

    def test_finance_grid_update_saves_months_and_audits(self):
        line = self.version.lines.get()
        data = {
            f"line-{line.pk}-{field}": "0"
            for field in MONTH_FIELDS
        }
        data[f"line-{line.pk}-m01"] = "1500.25"
        data[f"line-{line.pk}-m02"] = "249.75"
        response = self._client(self.finance).post(
            reverse("budget:version_grid_update", args=[self.version.pk]),
            data,
        )
        self.assertEqual(response.status_code, 302)
        line.refresh_from_db()
        self.assertEqual(line.m01, D("1500.25"))
        self.assertEqual(line.m02, D("249.75"))
        self.assertEqual(line.annual_amount, D("1750.00"))
        self.assertTrue(AuditLog.objects.filter(
            action="entity_updated", entity_type="budgetline",
            entity_id=str(line.pk),
        ).exists())

    def test_grid_update_is_atomic_when_a_value_is_invalid(self):
        line = self.version.lines.get()
        data = {
            f"line-{line.pk}-{field}": "0"
            for field in MONTH_FIELDS
        }
        data[f"line-{line.pk}-m01"] = "-1"
        response = self._client(self.finance).post(
            reverse("budget:version_grid_update", args=[self.version.pk]),
            data,
        )
        self.assertEqual(response.status_code, 400)
        line.refresh_from_db()
        self.assertEqual(line.m01, D("1000"))
        self.assertEqual(line.annual_amount, D("1000"))

    def test_auditor_cannot_use_grid_update(self):
        line = self.version.lines.get()
        data = {
            f"line-{line.pk}-{field}": "0"
            for field in MONTH_FIELDS
        }
        response = self._client(self.auditor).post(
            reverse("budget:version_grid_update", args=[self.version.pk]),
            data,
        )
        self.assertEqual(response.status_code, 403)

    def test_approved_version_rejects_grid_update(self):
        self.version.status = "approved"
        self.version.save()
        line = self.version.lines.get()
        data = {
            f"line-{line.pk}-{field}": "0"
            for field in MONTH_FIELDS
        }
        response = self._client(self.admin).post(
            reverse("budget:version_grid_update", args=[self.version.pk]),
            data,
        )
        self.assertEqual(response.status_code, 403)
        line.refresh_from_db()
        self.assertEqual(line.annual_amount, D("1000"))

    def test_admin_approves_locks_and_audits(self):
        c = self._client(self.admin)
        response = c.post(reverse("budget:version_approve",
                                  args=[self.version.pk]))
        self.assertEqual(response.status_code, 302)
        self.version.refresh_from_db()
        self.assertEqual(self.version.status, "approved")
        self.assertEqual(self.version.approved_by_id, self.admin.pk)
        entry = AuditLog.objects.filter(action="budget_approved",
                                        entity_id=str(self.version.pk)).first()
        self.assertIsNotNone(entry)
        self.assertIn("1000", entry.diff.get("annual_total", ""))

    def test_approved_version_rejects_line_edits_403(self):
        self.version.status = "approved"
        self.version.save()
        c = self._client(self.admin)
        line = self.version.lines.first()
        # GET edit form → 403 (locked)
        response = c.get(reverse("budget:line_edit", args=[line.pk]))
        self.assertEqual(response.status_code, 403)
        # POST delete → 403 as well
        response = c.post(reverse("budget:line_delete", args=[line.pk]))
        self.assertEqual(response.status_code, 403)
        self.assertTrue(BudgetLine.objects.filter(pk=line.pk).exists())

    def test_unapprove_reopens_with_audit(self):
        self.version.status = "approved"
        self.version.save()
        c = self._client(self.admin)
        c.post(reverse("budget:version_unapprove", args=[self.version.pk]))
        self.version.refresh_from_db()
        self.assertEqual(self.version.status, "draft")
        self.assertIsNone(self.version.approved_by)
        self.assertTrue(AuditLog.objects.filter(
            action="budget_unapproved",
            entity_id=str(self.version.pk)).exists())

    def test_new_revision_copies_lines_and_becomes_analysis_after_approval(self):
        c = self._client(self.admin)
        # approve v1 first → analysis = v1
        c.post(reverse("budget:version_approve", args=[self.version.pk]))
        self.assertEqual(services.get_analysis_version(self.budget).version, 1)
        # create revision
        response = c.post(reverse("budget:version_new_revision",
                                  args=[self.budget.pk]))
        self.assertEqual(response.status_code, 302)
        v2 = self.budget.versions.get(version=2)
        self.assertEqual(v2.status, "draft")
        self.assertEqual(v2.lines.count(), 1)
        self.assertEqual(v2.lines.first().annual_amount, D("1000"))
        # analysis still v1 (approved) until v2 approved
        self.assertEqual(services.get_analysis_version(self.budget).version, 1)
        c.post(reverse("budget:version_approve", args=[v2.pk]))
        self.assertEqual(services.get_analysis_version(self.budget).version, 2)
        self.assertTrue(AuditLog.objects.filter(
            action="budget_version_created",
            entity_id=str(v2.pk)).exists())

    def test_cannot_delete_budget_with_approved_version(self):
        self.version.status = "approved"
        self.version.save()
        c = self._client(self.admin)
        response = c.post(reverse("budget:budget_delete", args=[self.budget.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Budget.objects.filter(pk=self.budget.pk).exists())

    def test_can_delete_draft_only_budget_with_audit(self):
        c = self._client(self.admin)
        pk = self.budget.pk
        response = c.post(reverse("budget:budget_delete", args=[pk]))
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Budget.objects.filter(pk=pk).exists())
        self.assertTrue(AuditLog.objects.filter(
            action="entity_deleted", entity_type="budget",
            entity_id=str(pk)).exists())

    def test_approve_empty_version_refused(self):
        empty = make_version(self.budget, version=7)
        c = self._client(self.admin)
        response = c.post(reverse("budget:version_approve", args=[empty.pk]))
        self.assertEqual(response.status_code, 302)  # redirect back
        empty.refresh_from_db()
        self.assertEqual(empty.status, "draft")  # NOT approved

    def test_summary_page_shows_db_totals(self):
        c = self._client(self.admin)
        response = c.get(reverse("budget:version_summary",
                                 args=[self.version.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "1000")          # annual total
        self.assertContains(response, "النسخة المستخدمة في التحليل")

    def test_list_search_and_status_filter(self):
        fy25 = make_fy(code="FY2025", year=2025)
        old = make_budget(fy25, name="ميزانية السنة القديمة")
        make_version(old)  # draft version
        c = self._client(self.admin)
        url = reverse("budget:budget_list")
        self.assertEqual(c.get(url, {"q": "قديمة"}).context["total"], 1)
        self.assertEqual(c.get(url, {"q": "FIN"}).context["total"], 0)
        # status filter: budget whose version is draft
        self.assertEqual(c.get(url, {"status": "draft"}).context["total"], 2)
        self.assertEqual(c.get(url, {"status": "approved"}).context["total"], 0)
