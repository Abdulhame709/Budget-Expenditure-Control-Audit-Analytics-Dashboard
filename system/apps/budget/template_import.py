from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from django.db import transaction
from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.reference.models import Account, Currency

from .models import (
    BudgetTemplate,
    BudgetTemplateCell,
    BudgetTemplateColumn,
    BudgetTemplateRow,
    BudgetTemplateSheet,
)


MONTH_NAMES = {
    "يناير": 1, "فبراير": 2, "مارس": 3, "أبريل": 4, "ابريل": 4,
    "مايو": 5, "يونيو": 6, "يوليو": 7, "أغسطس": 8, "اغسطس": 8,
    "سبتمبر": 9, "أكتوبر": 10, "اكتوبر": 10, "نوفمبر": 11, "ديسمبر": 12,
}


def _text(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def _color(value):
    if not value:
        return ""
    rgb = getattr(value, "rgb", None)
    return rgb if isinstance(rgb, str) else ""


def _style(cell) -> dict:
    return {
        "bold": bool(cell.font.bold),
        "italic": bool(cell.font.italic),
        "font_color": _color(cell.font.color),
        "fill_color": _color(cell.fill.fgColor),
        "horizontal": cell.alignment.horizontal or "",
        "vertical": cell.alignment.vertical or "",
        "wrap_text": bool(cell.alignment.wrap_text),
        "border": bool(cell.border.left.style or cell.border.right.style
                       or cell.border.top.style or cell.border.bottom.style),
    }


def _sheet_purpose(name: str, values: list[str]) -> str:
    text = " ".join([name, *values[:80]])
    if "إجمالي" in text or "اجمالي" in text:
        return BudgetTemplateSheet.PURPOSE_SUMMARY
    if any(month in text for month in MONTH_NAMES) or "شهري" in text:
        return BudgetTemplateSheet.PURPOSE_MONTHLY
    if "تفصيلي" in text or "تفصيل" in text:
        return BudgetTemplateSheet.PURPOSE_DETAIL
    return BudgetTemplateSheet.PURPOSE_SUPPORTING


def _row_type(values: list[str], bold_count: int, merged: bool) -> str:
    text = " ".join(values).strip()
    if not text:
        return BudgetTemplateRow.TYPE_NOTE
    if "الإجمالي" in text or "الاجمالي" in text or text.startswith("إجمالي"):
        return BudgetTemplateRow.TYPE_TOTAL
    if "المجموع" in text or "جملة" in text:
        return BudgetTemplateRow.TYPE_SUBTOTAL
    if merged and len(values) <= 2:
        return BudgetTemplateRow.TYPE_TITLE
    if bold_count >= max(1, len(values) // 2) and not any(char.isdigit() for char in text):
        return BudgetTemplateRow.TYPE_HEADER
    return BudgetTemplateRow.TYPE_DATA


def _column_role(title: str) -> str:
    normalized = (title or "").strip().lower()
    if "رقم الحساب" in normalized or "كود الحساب" in normalized:
        return BudgetTemplateColumn.ROLE_ACCOUNT_CODE
    if "اسم الحساب" in normalized or normalized == "الحساب":
        return BudgetTemplateColumn.ROLE_ACCOUNT_NAME
    if "إجمالي" in normalized or "اجمالي" in normalized or "سنوي" in normalized:
        return BudgetTemplateColumn.ROLE_ANNUAL
    if "ملاحظ" in normalized or "بيان" == normalized:
        return BudgetTemplateColumn.ROLE_NOTES
    for month_name, month in MONTH_NAMES.items():
        if month_name in normalized:
            return f"month_{month:02d}"
    return BudgetTemplateColumn.ROLE_GENERIC


def _account_match(values: list[str], accounts_by_code: dict, accounts_by_name: dict):
    for value in values:
        cleaned = value.strip()
        if cleaned in accounts_by_code:
            return accounts_by_code[cleaned]
    joined = " ".join(values).strip().lower()
    for name, account in accounts_by_name.items():
        if name and (joined == name or name in joined):
            return account
    return None


@transaction.atomic
def import_budget_workbook(file_or_path, *, name: str, description: str = "", user=None):
    source_reference = getattr(file_or_path, "name", None) or str(file_or_path)
    keep_vba = str(source_reference).lower().endswith(".xlsm")
    workbook = load_workbook(file_or_path, data_only=False, keep_vba=keep_vba)
    source_name = Path(source_reference).name
    base_currency = Currency.objects.filter(is_base=True).first()
    template = BudgetTemplate.objects.create(
        name=name,
        description=description,
        source_filename=source_name,
        currency=base_currency,
        workbook_metadata={
            "sheet_count": len(workbook.sheetnames),
            "sheet_names": workbook.sheetnames,
            "calculation_mode": getattr(workbook.calculation, "calcMode", None),
        },
        created_by=user,
        updated_by=user,
    )
    accounts = list(Account.objects.filter(account_type="expense", is_active=True))
    accounts_by_code = {str(account.code).strip(): account for account in accounts}
    accounts_by_name = {account.name.strip().lower(): account for account in accounts}

    for position, worksheet in enumerate(workbook.worksheets, start=1):
        merged_ranges = [str(value) for value in worksheet.merged_cells.ranges]
        non_empty_text = [
            _text(cell.value)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        ]
        sheet = BudgetTemplateSheet.objects.create(
            template=template,
            name=worksheet.title,
            position=position,
            purpose=_sheet_purpose(worksheet.title, non_empty_text),
            state=worksheet.sheet_state,
            max_row=worksheet.max_row,
            max_column=worksheet.max_column,
            freeze_panes=_text(worksheet.freeze_panes),
            merged_ranges=merged_ranges,
            layout_metadata={
                "sheet_view": {
                    "show_grid_lines": worksheet.sheet_view.showGridLines,
                    "right_to_left": worksheet.sheet_view.rightToLeft,
                },
                "page_orientation": worksheet.page_setup.orientation,
            },
            created_by=user,
            updated_by=user,
        )

        column_objects = []
        for column_index in range(1, worksheet.max_column + 1):
            column_letter = get_column_letter(column_index)
            candidates = []
            for row_number in range(1, min(worksheet.max_row, 20) + 1):
                value = worksheet.cell(row_number, column_index).value
                if value not in (None, ""):
                    candidates.append(_text(value))
            title = candidates[0][:255] if candidates else ""
            dimension = worksheet.column_dimensions[column_letter]
            column_objects.append(BudgetTemplateColumn(
                sheet=sheet,
                column_index=column_index,
                column_letter=column_letter,
                title=title,
                role=_column_role(" ".join(candidates)),
                width=dimension.width,
                is_hidden=bool(dimension.hidden),
            ))
        BudgetTemplateColumn.objects.bulk_create(column_objects)
        columns = {column.column_index: column for column in column_objects}

        merged_rows = set()
        for merged_range in worksheet.merged_cells.ranges:
            merged_rows.update(range(merged_range.min_row, merged_range.max_row + 1))

        row_objects = []
        row_cell_specs = []
        for row_number in range(1, worksheet.max_row + 1):
            excel_cells = list(worksheet[row_number])
            populated = [cell for cell in excel_cells if cell.value not in (None, "")]
            if not populated and row_number not in merged_rows:
                continue
            values = [_text(cell.value) for cell in populated]
            row_dimension = worksheet.row_dimensions[row_number]
            matched_account = _account_match(values, accounts_by_code, accounts_by_name)
            main_account = None
            analytical_account = None
            if matched_account:
                if matched_account.level == 5:
                    main_account = matched_account
                elif matched_account.level == 6:
                    analytical_account = matched_account
                    main_account = matched_account.parent
            row_type = _row_type(
                values,
                sum(1 for cell in populated if cell.font.bold),
                row_number in merged_rows,
            )
            row = BudgetTemplateRow(
                sheet=sheet,
                row_number=row_number,
                position=row_number,
                row_type=row_type,
                label=" | ".join(values[:4])[:500],
                account=matched_account,
                main_account=main_account,
                analytical_account=analytical_account,
                is_included=bool(populated),
                height=row_dimension.height,
                is_hidden=bool(row_dimension.hidden),
            )
            row_objects.append(row)
            row_cell_specs.append((row, row_type, excel_cells))
        BudgetTemplateRow.objects.bulk_create(row_objects)

        cell_objects = []
        for row, row_type, excel_cells in row_cell_specs:
            for cell in excel_cells:
                raw_value = _text(cell.value)
                formula = raw_value if cell.data_type == "f" or raw_value.startswith("=") else ""
                cell_objects.append(BudgetTemplateCell(
                    row=row,
                    column=columns[cell.column],
                    coordinate=cell.coordinate,
                    raw_value=raw_value,
                    formula=formula,
                    is_editable=(
                        row_type == BudgetTemplateRow.TYPE_DATA
                        and not formula
                        and not isinstance(cell, MergedCell)
                    ),
                    data_type=cell.data_type or "",
                    number_format=cell.number_format or "",
                    style_metadata=_style(cell),
                ))
        BudgetTemplateCell.objects.bulk_create(cell_objects, batch_size=1000)
    workbook.close()
    return template


@transaction.atomic
def refresh_budget_workbook(template, file_or_path, *, user=None):
    """Import a newer workbook into the same template while preserving row links."""
    source_reference = getattr(file_or_path, "name", None) or str(file_or_path)
    temporary = import_budget_workbook(
        file_or_path,
        name=f"{template.name} — تحديث مؤقت",
        description="temporary refresh",
        user=user,
    )
    stats = {"sheets": 0, "rows": 0, "cells": 0, "new_sheets": 0}
    existing_sheets = {sheet.name: sheet for sheet in template.sheets.all()}
    next_position = template.sheets.order_by("-position").values_list(
        "position", flat=True,
    ).first() or 0

    for source_sheet in temporary.sheets.all().prefetch_related(
        "columns", "rows__cells",
    ):
        target_sheet = existing_sheets.get(source_sheet.name)
        if target_sheet is None:
            next_position += 1
            target_sheet = BudgetTemplateSheet.objects.create(
                template=template,
                name=source_sheet.name,
                position=next_position,
                purpose=source_sheet.purpose,
                state=source_sheet.state,
                max_row=source_sheet.max_row,
                max_column=source_sheet.max_column,
                freeze_panes=source_sheet.freeze_panes,
                merged_ranges=source_sheet.merged_ranges,
                layout_metadata=source_sheet.layout_metadata,
                created_by=user,
                updated_by=user,
            )
            stats["new_sheets"] += 1
        else:
            target_sheet.state = source_sheet.state
            target_sheet.max_row = source_sheet.max_row
            target_sheet.max_column = source_sheet.max_column
            target_sheet.freeze_panes = source_sheet.freeze_panes
            target_sheet.merged_ranges = source_sheet.merged_ranges
            target_sheet.layout_metadata = source_sheet.layout_metadata
            target_sheet.updated_by = user
            target_sheet.save()
        stats["sheets"] += 1

        target_columns = {column.column_index: column for column in target_sheet.columns.all()}
        column_map = {}
        for source_column in source_sheet.columns.all():
            target_column = target_columns.get(source_column.column_index)
            if target_column is None:
                target_column = BudgetTemplateColumn.objects.create(
                    sheet=target_sheet,
                    column_index=source_column.column_index,
                    column_letter=source_column.column_letter,
                    title=source_column.title,
                    role=source_column.role,
                    width=source_column.width,
                    is_hidden=source_column.is_hidden,
                )
            else:
                target_column.column_letter = source_column.column_letter
                target_column.title = source_column.title
                target_column.width = source_column.width
                target_column.is_hidden = source_column.is_hidden
                target_column.save()
            column_map[source_column.pk] = target_column

        target_rows = {row.row_number: row for row in target_sheet.rows.all()}
        for source_row in source_sheet.rows.all().prefetch_related("cells"):
            target_row = target_rows.get(source_row.row_number)
            if target_row is None:
                target_row = BudgetTemplateRow.objects.create(
                    sheet=target_sheet,
                    row_number=source_row.row_number,
                    position=source_row.position or source_row.row_number,
                    row_type=source_row.row_type,
                    label=source_row.label,
                    account=source_row.account,
                    main_account=source_row.main_account,
                    analytical_account=source_row.analytical_account,
                    is_included=source_row.is_included,
                    height=source_row.height,
                    is_hidden=source_row.is_hidden,
                )
            else:
                target_row.label = source_row.label
                target_row.height = source_row.height
                target_row.is_hidden = source_row.is_hidden
                target_row.save(update_fields=["label", "height", "is_hidden"])
            stats["rows"] += 1
            target_cells = {cell.column_id: cell for cell in target_row.cells.all()}
            for source_cell in source_row.cells.all():
                target_column = column_map[source_cell.column_id]
                target_cell = target_cells.get(target_column.pk)
                defaults = {
                    "coordinate": f"{target_column.column_letter}{target_row.row_number}",
                    "raw_value": source_cell.raw_value,
                    "override_value": None,
                    "formula": source_cell.formula,
                    "is_editable": source_cell.is_editable,
                    "data_type": source_cell.data_type,
                    "number_format": source_cell.number_format,
                    "style_metadata": source_cell.style_metadata,
                }
                if target_cell is None:
                    BudgetTemplateCell.objects.create(
                        row=target_row, column=target_column, **defaults,
                    )
                else:
                    for field, value in defaults.items():
                        setattr(target_cell, field, value)
                    target_cell.save()
                stats["cells"] += 1

    template.source_filename = Path(source_reference).name
    template.workbook_metadata = temporary.workbook_metadata
    template.updated_by = user
    template.save(update_fields=[
        "source_filename", "workbook_metadata", "updated_by", "updated_at",
    ])
    temporary.delete()
    return stats
