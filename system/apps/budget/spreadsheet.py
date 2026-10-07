"""Univer workbook adapter for editable budget-template sheets."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl.utils import get_column_letter, range_boundaries

from .models import (
    BudgetTemplateCell,
    BudgetTemplateColumn,
    BudgetTemplateRow,
    BudgetTemplateSheet,
)


class SpreadsheetPayloadError(ValueError):
    pass


def _number(value):
    if value in (None, ""):
        return value
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return value
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def _cell_value(cell: BudgetTemplateCell):
    value = cell.effective_value
    if cell.data_type in {"n", "number", "numeric"}:
        return _number(value)
    return value


def _univer_column_width(width) -> float:
    if width is None:
        return 120
    value = float(width)
    if value < 45:
        value = (value * 7) + 12
    return max(value, 120)


def _univer_style(cell: BudgetTemplateCell) -> dict:
    metadata = cell.style_metadata or {}
    if isinstance(metadata.get("univer"), dict):
        return metadata["univer"]
    style = {}
    if metadata.get("bold"):
        style["bl"] = 1
    if metadata.get("italic"):
        style["it"] = 1
    if metadata.get("wrap_text"):
        style["tb"] = 1
    horizontal = metadata.get("horizontal")
    if horizontal in {"left", "center", "right"}:
        style["ht"] = {"left": 1, "center": 2, "right": 3}[horizontal]
    return style


def workbook_snapshot(sheet: BudgetTemplateSheet) -> dict:
    columns = list(sheet.columns.order_by("column_index"))
    rows = list(
        sheet.rows.order_by("position", "row_number", "pk").prefetch_related("cells")
    )
    column_positions = {column.pk: index for index, column in enumerate(columns)}
    row_number_positions = {row.row_number: index for index, row in enumerate(rows)}
    column_number_positions = {
        column.column_index: index for index, column in enumerate(columns)
    }
    cell_data = {}
    row_data = {}
    column_data = {}

    for index, column in enumerate(columns):
        column_data[index] = {
            "w": _univer_column_width(column.width),
            "hd": bool(column.is_hidden),
            "custom": {"budgetColumnId": column.pk},
        }

    for row_index, row in enumerate(rows):
        row_data[row_index] = {
            "h": float(row.height) if row.height is not None else 24,
            "hd": bool(row.is_hidden),
            "custom": {
                "budgetRowId": row.pk,
                "budgetRowType": row.row_type,
            },
        }
        rendered = {}
        for cell in row.cells.all():
            column_index = column_positions.get(cell.column_id)
            if column_index is None:
                continue
            payload = {
                "v": _cell_value(cell),
                "custom": {
                    "budgetCellId": cell.pk,
                    "budgetRowId": row.pk,
                    "budgetColumnId": cell.column_id,
                },
            }
            if cell.formula:
                payload["f"] = cell.formula
            style = _univer_style(cell)
            if style:
                payload["s"] = style
            rendered[column_index] = payload
        if rendered:
            cell_data[row_index] = rendered

    merges = []
    for value in sheet.merged_ranges or []:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(str(value))
        except (TypeError, ValueError):
            continue
        merged_rows = [
            row_number_positions[number]
            for number in range(min_row, max_row + 1)
            if number in row_number_positions
        ]
        merged_columns = [
            column_number_positions[number]
            for number in range(min_col, max_col + 1)
            if number in column_number_positions
        ]
        if merged_rows and merged_columns:
            merges.append({
                "startRow": min(merged_rows),
                "endRow": max(merged_rows),
                "startColumn": min(merged_columns),
                "endColumn": max(merged_columns),
            })

    worksheet_id = f"budget-sheet-{sheet.pk}"
    workbook_id = f"budget-workbook-{sheet.template_id}-{sheet.pk}"
    return {
        "id": workbook_id,
        "name": sheet.template.name,
        "locale": "arSA",
        "styles": {},
        "sheetOrder": [worksheet_id],
        "sheets": {
            worksheet_id: {
                "id": worksheet_id,
                "name": sheet.name,
                "rowCount": max(len(rows), 1),
                "columnCount": max(len(columns), 1),
                "cellData": cell_data,
                "rowData": row_data,
                "columnData": column_data,
                "mergeData": merges,
                "rightToLeft": 1,
                "showGridlines": 1,
            }
        },
    }


def _indexed(value) -> dict[int, dict]:
    if isinstance(value, list):
        return {index: item for index, item in enumerate(value) if isinstance(item, dict)}
    if not isinstance(value, dict):
        return {}
    output = {}
    for key, item in value.items():
        if not isinstance(item, dict):
            continue
        try:
            output[int(key)] = item
        except (TypeError, ValueError):
            continue
    return output


def _custom_id(payload, key):
    custom = payload.get("custom") if isinstance(payload, dict) else None
    if not isinstance(custom, dict):
        return None
    try:
        return int(custom.get(key))
    except (TypeError, ValueError):
        return None


def _dimension_limit(metadata, cell_data, merges, axis):
    indices = set(metadata)
    if axis == "row":
        indices.update(cell_data)
        for item in merges:
            if not isinstance(item, dict):
                continue
            try:
                indices.add(int(item.get("endRow", 0)))
            except (TypeError, ValueError):
                continue
    else:
        for cells in cell_data.values():
            indices.update(cells)
        for item in merges:
            if not isinstance(item, dict):
                continue
            try:
                indices.add(int(item.get("endColumn", 0)))
            except (TypeError, ValueError):
                continue
    return max(indices, default=0) + 1


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    return str(value)


def _data_type(value) -> str:
    return "n" if isinstance(value, (int, float)) and not isinstance(value, bool) else "s"


def _merge_ranges(merges, row_count, column_count):
    ranges = []
    for merge in merges:
        if not isinstance(merge, dict):
            continue
        try:
            start_row = max(0, int(merge.get("startRow", 0)))
            end_row = min(row_count - 1, int(merge.get("endRow", start_row)))
            start_column = max(0, int(merge.get("startColumn", 0)))
            end_column = min(column_count - 1, int(merge.get("endColumn", start_column)))
        except (TypeError, ValueError):
            continue
        if end_row < start_row or end_column < start_column:
            continue
        ranges.append(
            f"{get_column_letter(start_column + 1)}{start_row + 1}:"
            f"{get_column_letter(end_column + 1)}{end_row + 1}"
        )
    return ranges


@transaction.atomic
def save_workbook_snapshot(sheet: BudgetTemplateSheet, workbook: dict) -> dict:
    if not isinstance(workbook, dict):
        raise SpreadsheetPayloadError("بيانات المصنف غير صالحة.")
    worksheets = workbook.get("sheets")
    if not isinstance(worksheets, dict) or not worksheets:
        raise SpreadsheetPayloadError("لا توجد ورقة قابلة للحفظ.")
    worksheet = next(iter(worksheets.values()))
    if not isinstance(worksheet, dict):
        raise SpreadsheetPayloadError("بيانات الورقة غير صالحة.")

    cell_data = {
        row_index: _indexed(cells)
        for row_index, cells in _indexed(worksheet.get("cellData")).items()
    }
    row_data = _indexed(worksheet.get("rowData"))
    column_data = _indexed(worksheet.get("columnData"))
    merges = worksheet.get("mergeData") or []
    if not isinstance(merges, list):
        merges = []
    row_count = _dimension_limit(row_data, cell_data, merges, "row")
    column_count = _dimension_limit(column_data, cell_data, merges, "column")
    if row_count > 5000 or column_count > 250:
        raise SpreadsheetPayloadError("حجم الورقة يتجاوز الحد المسموح للتجربة.")

    existing_rows = {row.pk: row for row in sheet.rows.select_for_update()}
    existing_columns = {column.pk: column for column in sheet.columns.select_for_update()}
    ordered_rows = []
    ordered_columns = []
    used_row_ids = set()
    used_column_ids = set()

    row_temp = max([row.row_number for row in existing_rows.values()] or [0]) + row_count + 1000
    for offset, row in enumerate(existing_rows.values()):
        row.row_number = row_temp + offset
        row.position = row_temp + offset
        row.save(update_fields=["row_number", "position"])
    column_temp = max([column.column_index for column in existing_columns.values()] or [0]) + column_count + 1000
    for offset, column in enumerate(existing_columns.values()):
        column.column_index = column_temp + offset
        column.column_letter = get_column_letter(column_temp + offset)
        column.save(update_fields=["column_index", "column_letter"])

    for index in range(row_count):
        metadata = row_data.get(index, {})
        row_id = _custom_id(metadata, "budgetRowId")
        row = existing_rows.get(row_id) if row_id not in used_row_ids else None
        if row is None:
            row = BudgetTemplateRow(sheet=sheet, row_type=BudgetTemplateRow.TYPE_DATA)
        else:
            used_row_ids.add(row.pk)
        row.row_number = index + 1
        row.position = index + 1
        row.height = metadata.get("h") or None
        row.is_hidden = bool(metadata.get("hd", False))
        row.save()
        ordered_rows.append(row)

    for index in range(column_count):
        metadata = column_data.get(index, {})
        column_id = _custom_id(metadata, "budgetColumnId")
        column = existing_columns.get(column_id) if column_id not in used_column_ids else None
        if column is None:
            column = BudgetTemplateColumn(sheet=sheet)
        else:
            used_column_ids.add(column.pk)
        column.column_index = index + 1
        column.column_letter = get_column_letter(index + 1)
        column.width = metadata.get("w") or None
        column.is_hidden = bool(metadata.get("hd", False))
        column.save()
        ordered_columns.append(column)

    BudgetTemplateRow.objects.filter(sheet=sheet).exclude(pk__in=[row.pk for row in ordered_rows]).delete()
    BudgetTemplateColumn.objects.filter(sheet=sheet).exclude(pk__in=[column.pk for column in ordered_columns]).delete()

    existing_cells = {
        (cell.row_id, cell.column_id): cell
        for cell in BudgetTemplateCell.objects.filter(row__sheet=sheet).select_for_update()
    }
    kept_cell_ids = []
    for row_index, cells in cell_data.items():
        if row_index >= len(ordered_rows):
            continue
        for column_index, payload in cells.items():
            if column_index >= len(ordered_columns):
                continue
            has_content = any(key in payload for key in ("v", "f", "s", "custom"))
            if not has_content:
                continue
            row = ordered_rows[row_index]
            column = ordered_columns[column_index]
            cell = existing_cells.get((row.pk, column.pk))
            value = payload.get("v")
            formula = _text(payload.get("f"))
            style = payload.get("s") if isinstance(payload.get("s"), dict) else {}
            if cell is None:
                cell = BudgetTemplateCell(row=row, column=column, is_editable=True)
                cell.raw_value = _text(value)
                cell.override_value = None
            elif formula:
                cell.raw_value = _text(value)
                cell.override_value = None
            else:
                cell.override_value = None if _text(value) == cell.raw_value else _text(value)
            cell.coordinate = f"{column.column_letter}{row.row_number}"
            cell.formula = formula
            cell.data_type = _data_type(value)
            metadata = dict(cell.style_metadata or {})
            if style:
                metadata["univer"] = style
            else:
                metadata.pop("univer", None)
            cell.style_metadata = metadata
            cell.save()
            kept_cell_ids.append(cell.pk)

    BudgetTemplateCell.objects.filter(row__sheet=sheet).exclude(pk__in=kept_cell_ids).delete()
    sheet.max_row = row_count
    sheet.max_column = column_count
    sheet.merged_ranges = _merge_ranges(merges, row_count, column_count)
    sheet.save(update_fields=["max_row", "max_column", "merged_ranges", "updated_at"])
    return {"rows": row_count, "columns": column_count, "cells": len(kept_cell_ids)}
