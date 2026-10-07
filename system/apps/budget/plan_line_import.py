"""Excel import helpers for the section-level bulk budget editor."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from apps.reference.models import Account

from .models import BudgetPlanLine


HEADER_ALIASES = {
    "position": {"الترتيب", "رقم", "م", "position"},
    "row_type": {"نوع الصف", "النوع", "row type", "row_type"},
    "display_name": {"اسم البند", "البيان", "اسم العرض", "البند", "display name", "display_name"},
    "main_account": {"الحساب الرئيسي", "رمز الحساب الرئيسي", "main account", "main_account"},
    "analytical_account": {"الحساب التحليلي", "رمز الحساب التحليلي", "analytical account", "analytical_account"},
    "input_mode": {"طريقة الإدخال", "طريقة الادخال", "input mode", "input_mode"},
    "distribution_method": {"طريقة التوزيع", "التوزيع", "distribution", "distribution_method"},
    "annual_amount": {"المبلغ السنوي", "الإجمالي السنوي", "الاجمالي السنوي", "المبلغ", "annual amount", "annual_amount"},
    "estimation_basis": {"أساس التقدير", "الملاحظات", "ملاحظات", "notes", "estimation_basis"},
    "is_included": {"مضمن", "مضمّن", "يدخل في الإجمالي", "is included", "is_included"},
}

ROW_TYPE_ALIASES = {
    "عنوان": BudgetPlanLine.TYPE_TITLE,
    "حساب رئيسي": BudgetPlanLine.TYPE_MAIN_ACCOUNT,
    "حساب تحليلي": BudgetPlanLine.TYPE_ANALYTICAL_ACCOUNT,
    "تفصيل": BudgetPlanLine.TYPE_DETAIL,
    "بند": BudgetPlanLine.TYPE_DETAIL,
    "إجمالي فرعي": BudgetPlanLine.TYPE_SUBTOTAL,
    "اجمالي فرعي": BudgetPlanLine.TYPE_SUBTOTAL,
    "إجمالي عام": BudgetPlanLine.TYPE_TOTAL,
    "اجمالي عام": BudgetPlanLine.TYPE_TOTAL,
    "ملاحظة": BudgetPlanLine.TYPE_NOTE,
}

INPUT_MODE_ALIASES = {
    "بدون إدخال": BudgetPlanLine.INPUT_NONE,
    "بدون ادخال": BudgetPlanLine.INPUT_NONE,
    "مبلغ سنوي": BudgetPlanLine.INPUT_ANNUAL,
    "سنوي": BudgetPlanLine.INPUT_ANNUAL,
}

DISTRIBUTION_ALIASES = {
    "بدون توزيع": BudgetPlanLine.DIST_NONE,
    "متساوي": BudgetPlanLine.DIST_EQUAL,
    "متساو": BudgetPlanLine.DIST_EQUAL,
    "متساوٍ على 12 شهرًا": BudgetPlanLine.DIST_EQUAL,
    "على 12 شهر": BudgetPlanLine.DIST_EQUAL,
}


class PlanLineImportError(ValueError):
    pass


def _text(value):
    return " ".join(str(value or "").strip().split())


def _normalized(value):
    return _text(value).casefold().replace("ـ", "")


def _header_key(value):
    normalized = _normalized(value)
    for key, aliases in HEADER_ALIASES.items():
        if normalized in {_normalized(alias) for alias in aliases}:
            return key
    return None


def _choice(value, choices, aliases, default):
    text = _text(value)
    if not text:
        return default
    valid = {key for key, _label in choices}
    if text in valid:
        return text
    normalized_aliases = {_normalized(label): key for label, key in aliases.items()}
    return normalized_aliases.get(_normalized(text), default)


def _amount(value, row_number, errors):
    if value in (None, ""):
        return Decimal("0")
    try:
        amount = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        errors.append(f"الصف {row_number}: المبلغ السنوي غير صالح.")
        return Decimal("0")
    if amount < 0:
        errors.append(f"الصف {row_number}: لا يمكن استيراد مبلغ سالب.")
        return Decimal("0")
    return amount


def _included(value):
    if value in (None, ""):
        return True
    return _normalized(value) not in {"0", "لا", "false", "no", "غير مضمن", "غير مضمّن"}


def _account(value, *, level, parent=None):
    text = _text(value)
    if not text:
        return None
    queryset = Account.objects.filter(
        account_type="expense", is_active=True, level=level,
    )
    if parent is not None:
        queryset = queryset.filter(parent=parent)
    account = queryset.filter(code=text).first()
    if account:
        return account
    return queryset.filter(name__iexact=text).first()


def parse_plan_lines_workbook(file_object, section):
    """Return editable form initial data plus row-level warnings.

    If a workbook has a sheet named after the selected section it is used;
    otherwise the active sheet is imported.
    """
    try:
        workbook = load_workbook(file_object, data_only=True, read_only=True)
    except Exception as exc:
        raise PlanLineImportError("تعذر قراءة ملف Excel. تأكد أنه بصيغة xlsx سليمة.") from exc
    try:
        worksheet = workbook.active
        section_name = _normalized(section.name)
        for candidate in workbook.worksheets:
            if _normalized(candidate.title) == section_name:
                worksheet = candidate
                break

        header_row = None
        columns = {}
        for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True), start=1):
            mapped = {index: _header_key(value) for index, value in enumerate(row)}
            mapped = {index: key for index, key in mapped.items() if key}
            if "display_name" in mapped.values():
                header_row = row_number
                columns = mapped
                break
        if header_row is None:
            raise PlanLineImportError(
                "لم يتم العثور على صف عناوين يحتوي على عمود «اسم البند» أو «البيان»."
            )

        rows = []
        errors = []
        next_position = section.lines.order_by("-position").values_list("position", flat=True).first() or 0
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            payload = {key: values[index] if index < len(values) else None for index, key in columns.items()}
            display_name = _text(payload.get("display_name"))
            if not display_name and not any(value not in (None, "") for value in values):
                continue
            if not display_name:
                errors.append(f"الصف {row_number}: تم تجاهله لعدم وجود اسم بند.")
                continue

            requested_position = payload.get("position")
            try:
                position = int(requested_position) if requested_position not in (None, "") else next_position + 1
            except (TypeError, ValueError):
                position = next_position + 1
                errors.append(f"الصف {row_number}: الترتيب غير صالح؛ تم اقتراح {position}.")
            next_position = max(next_position, position)
            row_type = _choice(
                payload.get("row_type"), BudgetPlanLine.ROW_TYPE_CHOICES,
                ROW_TYPE_ALIASES, BudgetPlanLine.TYPE_DETAIL,
            )
            main_account = _account(payload.get("main_account"), level=5)
            analytical_account = _account(
                payload.get("analytical_account"), level=6, parent=main_account,
            )
            if payload.get("main_account") and not main_account:
                errors.append(f"الصف {row_number}: لم تتم مطابقة الحساب الرئيسي.")
            if payload.get("analytical_account") and not analytical_account:
                errors.append(f"الصف {row_number}: لم تتم مطابقة الحساب التحليلي مع الحساب الرئيسي.")
            input_mode = _choice(
                payload.get("input_mode"), BudgetPlanLine.INPUT_MODE_CHOICES,
                INPUT_MODE_ALIASES, BudgetPlanLine.INPUT_ANNUAL,
            )
            distribution = _choice(
                payload.get("distribution_method"), BudgetPlanLine.DISTRIBUTION_CHOICES,
                DISTRIBUTION_ALIASES,
                BudgetPlanLine.DIST_EQUAL if input_mode == BudgetPlanLine.INPUT_ANNUAL else BudgetPlanLine.DIST_NONE,
            )
            rows.append({
                "position": position,
                "row_type": row_type,
                "display_name": display_name,
                "main_account": main_account,
                "analytical_account": analytical_account,
                "input_mode": input_mode,
                "distribution_method": distribution,
                "annual_amount": _amount(payload.get("annual_amount"), row_number, errors),
                "estimation_basis": _text(payload.get("estimation_basis")),
                "is_included": _included(payload.get("is_included")),
            })
        if not rows:
            raise PlanLineImportError("لا توجد بنود قابلة للاستيراد في ورقة Excel المحددة.")
        return rows, errors, worksheet.title
    finally:
        workbook.close()
