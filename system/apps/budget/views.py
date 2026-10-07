"""PHASE 5 — Budget Management views.

Flow: Budget (one/FY) → BudgetVersion (draft ⇄ approved; revision = new
version copying lines) → BudgetLine (manual monthly entry).

Enforcement:
- page   : URL_PERMISSIONS (budget.view / budget.edit) via middleware
- action : @require_permission on every mutation
- object : APPROVED versions are locked (line edits need a new revision)
- numbers: every total comes from DB aggregates (services.py)
"""
from __future__ import annotations

from django.contrib import messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db import models, transaction
from django.db.models import Count, Q
from django.db.models.deletion import ProtectedError
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST

from apps.accounts.permissions import has_perm, require_permission
from apps.governance.models import AuditLog
from apps.reference.services import changes_between, log_action, snapshot

from . import forms, services
from .models import (
    MONTH_FIELDS,
    Budget,
    BudgetLine,
    BudgetPlan,
    BudgetPlanLine,
    BudgetPlanSection,
    BudgetTemplate,
    BudgetTemplateCell,
    BudgetTemplateColumn,
    BudgetTemplateRow,
    BudgetTemplateSheet,
    BudgetVersion,
)
from .template_import import import_budget_workbook, refresh_budget_workbook
from .plan_line_import import PlanLineImportError, parse_plan_lines_workbook
from . import template_runtime


# ---------------------------------------------------------------- list
def budget_list(request):
    q = (request.GET.get("q") or "").strip()
    status = (request.GET.get("status") or "").strip()
    fy = (request.GET.get("fy") or "").strip()

    qs = Budget.objects.select_related("fiscal_year")
    if q:
        qs = qs.filter(
            Q(name__icontains=q)
            | Q(fiscal_year__code__icontains=q)
            | Q(fiscal_year__name__icontains=q)
        )
    if status in dict(BudgetVersion.STATUS_CHOICES):
        qs = qs.filter(versions__status=status).distinct()
    if fy.isdigit():
        qs = qs.filter(fiscal_year_id=int(fy))

    paginator = Paginator(qs, 15)
    try:
        page = paginator.get_page(request.GET.get("page"))
    except PageNotAnInteger:
        page = paginator.get_page(1)
    except EmptyPage:
        page = paginator.get_page(paginator.num_pages)

    rows = []
    for b in page:
        latest = b.latest_version
        analysis = b.analysis_version
        rows.append({
            "obj": b,
            "latest": latest,
            "analysis": analysis,
            "url": reverse("budget:budget_detail", args=[b.pk]),
        })

    return render(request, "budget/budget_list.html", {
        "rows": rows, "q": q, "total": paginator.count, "page_obj": page,
        "active_status": status,
        "status_choices": BudgetVersion.STATUS_CHOICES,
        "can_edit": has_perm(request.user, "budget.edit"),
        "create_url": reverse("budget:budget_create"),
    })


# ---------------------------------------------------------------- header CRUD
@require_permission("budget.edit")
def budget_create(request):
    if request.method == "POST":
        form = forms.BudgetForm(request.POST)
        if form.is_valid():
            budget = form.save(commit=False)
            budget.created_by = request.user
            budget.updated_by = request.user
            budget.save()
            version = BudgetVersion.objects.create(
                budget=budget, version=1, status=BudgetVersion.STATUS_DRAFT,
            )
            log_action(action="entity_created", entity_type="budget",
                       entity_id=budget.pk, diff=snapshot(budget), request=request)
            log_action(action="budget_version_created", entity_type="budgetversion",
                       entity_id=version.pk,
                       diff={"budget": str(budget), "version": 1,
                             "reason": "إنشاء الميزانية"},
                       request=request)
            messages.success(request, f"تم إنشاء الميزانية «{budget}» ونسخة المسودة v1.")
            return redirect("budget:budget_detail", pk=budget.pk)
    else:
        form = forms.BudgetForm()
    return render(request, "budget/budget_form.html", {
        "form": form, "is_create": True, "title": "ميزانية جديدة",
        "cancel_url": reverse("budget:budget_list"),
    })


def budget_detail(request, pk):
    budget = get_object_or_404(Budget.objects.select_related("fiscal_year"), pk=pk)
    versions = []
    for v in budget.versions.all():
        versions.append({
            "obj": v,
            "is_analysis": budget.analysis_version is not None
            and budget.analysis_version.pk == v.pk,
            "totals": services.version_totals(v),
        })
    return render(request, "budget/budget_detail.html", {
        "budget": budget, "versions": versions,
        "analysis": budget.analysis_version,
        "can_edit": has_perm(request.user, "budget.edit"),
        "line_form_url_base": None,
        "edit_url": reverse("budget:budget_edit", args=[budget.pk]),
        "delete_url": reverse("budget:budget_delete", args=[budget.pk]),
        "new_revision_url": reverse("budget:version_new_revision", args=[budget.pk])
        if budget.latest_version else "",
    })


@require_permission("budget.edit")
def budget_edit(request, pk):
    budget = get_object_or_404(Budget, pk=pk)
    before = snapshot(budget, ["name", "notes"])
    if request.method == "POST":
        form = forms.BudgetEditForm(request.POST, instance=budget)
        if form.is_valid():
            budget = form.save(commit=False)
            budget.updated_by = request.user
            budget.save()
            diff = changes_between(budget, before)
            if diff:
                log_action(action="entity_updated", entity_type="budget",
                           entity_id=budget.pk, diff=diff, request=request)
                messages.success(request, "تم تعديل الميزانية وتسجيل الفروق في سجل التدقيق.")
            else:
                messages.info(request, "لا توجد تغييرات لحفظها.")
            return redirect("budget:budget_detail", pk=budget.pk)
    else:
        form = forms.BudgetEditForm(instance=budget)
    return render(request, "budget/budget_form.html", {
        "form": form, "is_create": False, "title": f"تعديل: {budget}",
        "cancel_url": reverse("budget:budget_detail", args=[budget.pk]),
    })


@require_permission("budget.edit")
def budget_delete(request, pk):
    budget = get_object_or_404(Budget, pk=pk)
    if request.method != "POST":
        messages.error(request, "الحذف يتطلب تأكيدًا (POST).")
        return redirect("budget:budget_detail", pk=pk)
    if budget.versions.filter(status=BudgetVersion.STATUS_APPROVED).exists():
        messages.error(
            request,
            "لا يمكن حذف ميزانية تحتوي نسخة معتمدة — الاعتماد يعني أن الأرقام "
            "أصبحت أساس التحليل. (أنشئ نسخة مراجعة بدلًا من ذلك.)",
        )
        return redirect("budget:budget_detail", pk=pk)
    before = snapshot(budget, ["name", "notes", "fiscal_year"])
    label = str(budget)
    budget.delete()
    log_action(action="entity_deleted", entity_type="budget",
               entity_id=pk, diff={"snapshot": before, "label": label},
               request=request)
    messages.success(request, f"تم حذف الميزانية «{label}» وتسجيله في سجل التدقيق.")
    return redirect("budget:budget_list")


# ---------------------------------------------------------------- versions
def _version_detail_context(request, version, *, grid_forms=None):
    lines = list(version.lines.select_related(
        "department", "account", "analytical_account", "expense_category"))
    totals = services.version_totals(version)
    can_edit = has_perm(request.user, "budget.edit")
    locked = version.is_approved  # approved ⇒ lines locked
    if grid_forms is None and can_edit and not locked:
        grid_forms = [
            forms.BudgetLineGridForm(instance=line, prefix=f"line-{line.pk}")
            for line in lines
        ]
    form_by_line_id = {
        form.instance.pk: form for form in (grid_forms or [])
    }
    grid_rows = [
        {"line": line, "form": form_by_line_id.get(line.pk)}
        for line in lines
    ]
    return {
        "version": version, "budget": version.budget, "lines": lines,
        "grid_rows": grid_rows,
        "totals": totals, "can_edit": can_edit, "locked": locked,
        "is_analysis": version.is_analysis_version,
        "summary_url": reverse("budget:version_summary", args=[version.pk]),
        "approve_url": reverse("budget:version_approve", args=[version.pk]),
        "unapprove_url": reverse("budget:version_unapprove", args=[version.pk]),
        "edit_notes_url": reverse("budget:version_edit", args=[version.pk]),
        "line_create_url": reverse("budget:line_create", args=[version.pk]),
        "grid_update_url": reverse(
            "budget:version_grid_update", args=[version.pk]),
        "new_revision_url": reverse(
            "budget:version_new_revision", args=[version.budget_id]),
    }


def version_detail(request, pk):
    version = get_object_or_404(
        BudgetVersion.objects.select_related("budget__fiscal_year"), pk=pk)
    return render(
        request,
        "budget/version_detail.html",
        _version_detail_context(request, version),
    )


@require_permission("budget.edit")
@require_POST
def version_grid_update(request, pk):
    version = get_object_or_404(
        BudgetVersion.objects.select_related("budget__fiscal_year"), pk=pk)
    _require_draft(version, "التعديل الشبكي")
    lines = list(version.lines.select_related(
        "department", "account", "analytical_account", "expense_category"))
    if not lines:
        messages.info(request, "لا توجد سطور موازنة لتعديلها.")
        return redirect("budget:version_detail", pk=version.pk)

    before_by_line_id = {
        line.pk: snapshot(line, [*MONTH_FIELDS, "annual_amount"])
        for line in lines
    }
    grid_forms = [
        forms.BudgetLineGridForm(
            request.POST, instance=line, prefix=f"line-{line.pk}"
        )
        for line in lines
    ]
    validation_results = [form.is_valid() for form in grid_forms]
    if not all(validation_results):
        messages.error(
            request,
            "لم تُحفظ التغييرات: صحّح القيم المعلّمة داخل جدول الموازنة.",
        )
        return render(
            request,
            "budget/version_detail.html",
            _version_detail_context(request, version, grid_forms=grid_forms),
            status=400,
        )

    changed_count = 0
    with transaction.atomic():
        for form in grid_forms:
            line = form.instance
            before = before_by_line_id[line.pk]
            line = form.save(commit=False)
            line.updated_by = request.user
            line.save()
            diff = changes_between(line, before)
            if diff:
                changed_count += 1
                log_action(
                    action="entity_updated",
                    entity_type="budgetline",
                    entity_id=line.pk,
                    diff=diff,
                    request=request,
                )

    if changed_count:
        messages.success(
            request,
            f"تم حفظ {changed_count} من سطور الموازنة وتسجيل الفروق.",
        )
    else:
        messages.info(request, "لم توجد تغييرات جديدة للحفظ.")
    return redirect("budget:version_detail", pk=version.pk)


@require_permission("budget.edit")
def version_edit(request, pk):
    version = get_object_or_404(BudgetVersion, pk=pk)
    before = snapshot(version, ["notes"])
    if request.method == "POST":
        form = forms.BudgetVersionForm(request.POST, instance=version)
        if form.is_valid():
            version = form.save(commit=False)
            version.updated_by = request.user
            version.save()
            diff = changes_between(version, before)
            if diff:
                log_action(action="entity_updated", entity_type="budgetversion",
                           entity_id=version.pk, diff=diff, request=request)
                messages.success(request, "تم حفظ ملاحظات النسخة.")
            return redirect("budget:version_detail", pk=version.pk)
    else:
        form = forms.BudgetVersionForm(instance=version)
    return render(request, "budget/budget_form.html", {
        "form": form, "is_create": False,
        "title": f"ملاحظات النسخة: {version}",
        "cancel_url": reverse("budget:version_detail", args=[version.pk]),
    })


@require_permission("budget.edit")
def version_approve(request, pk):
    """draft → approved: locks the version and makes it eligible for analysis."""
    version = get_object_or_404(BudgetVersion, pk=pk)
    if request.method != "POST":
        messages.error(request, "الاعتماد يتطلب POST.")
        return redirect("budget:version_detail", pk=pk)
    if version.is_approved:
        messages.info(request, "النسخة معتمدة بالفعل.")
        return redirect("budget:version_detail", pk=pk)
    if not version.lines.exists():
        messages.error(request, "لا يمكن اعتماد نسخة بلا سطور ميزانية.")
        return redirect("budget:version_detail", pk=pk)
    version.status = BudgetVersion.STATUS_APPROVED
    version.approved_by = request.user
    version.approved_at = timezone.now()
    version.updated_by = request.user
    version.save(update_fields=["status", "approved_by", "approved_at",
                                "updated_by", "updated_at"])
    log_action(action="budget_approved", entity_type="budgetversion",
               entity_id=version.pk,
               diff={"version": str(version), "lines": version.lines.count(),
                     "annual_total": str(services.version_totals(version)["annual_total"])},
               request=request)
    messages.success(
        request,
        f"تم اعتماد النسخة {version} وأصبحت هي النسخة المستخدمة في التحليل — "
        "وأُقفلت للتعديل (التعديل عبر نسخة مراجعة جديدة).",
    )
    return redirect("budget:version_detail", pk=version.pk)


@require_permission("budget.edit")
def version_unapprove(request, pk):
    version = get_object_or_404(BudgetVersion, pk=pk)
    if request.method != "POST":
        messages.error(request, "الإجراء يتطلب POST.")
        return redirect("budget:version_detail", pk=pk)
    if not version.is_approved:
        messages.info(request, "النسخة ليست معتمدة.")
        return redirect("budget:version_detail", pk=pk)
    version.status = BudgetVersion.STATUS_DRAFT
    version.approved_by = None
    version.approved_at = None
    version.updated_by = request.user
    version.save(update_fields=["status", "approved_by", "approved_at",
                                "updated_by", "updated_at"])
    log_action(action="budget_unapproved", entity_type="budgetversion",
               entity_id=version.pk, diff={"version": str(version)},
               request=request)
    messages.warning(request, f"تم إلغاء اعتماد النسخة {version} وأصبحت مسودة قابلة للتعديل.")
    return redirect("budget:version_detail", pk=version.pk)


@require_permission("budget.edit")
def version_new_revision(request, budget_pk):
    """Revised budget: version N+1 (draft) copying all lines of the
    latest version — the approved one stays intact for analysis."""
    budget = get_object_or_404(Budget, pk=budget_pk)
    if request.method != "POST":
        messages.error(request, "الإجراء يتطلب POST.")
        return redirect("budget:budget_detail", pk=budget.pk)
    source = budget.latest_version
    if source is None:
        messages.error(request, "لا توجد نسخة لنسخها.")
        return redirect("budget:budget_detail", pk=budget.pk)
    new_version = BudgetVersion.objects.create(
        budget=budget, version=source.version + 1,
        status=BudgetVersion.STATUS_DRAFT, created_by=request.user,
        updated_by=request.user,
        notes=f"مراجعة للنسخة {source.version}" + (
            " (المعتمدة)" if source.is_approved else ""),
    )
    copied = 0
    for line in source.lines.all():
        BudgetLine.objects.create(
            version=new_version, department_id=line.department_id,
            account_id=line.account_id,
            analytical_account_id=line.analytical_account_id,
            expense_category_id=line.expense_category_id,
            **{f: getattr(line, f) for f in MONTH_FIELDS},
            annual_amount=line.annual_amount, notes=line.notes,
            created_by=request.user, updated_by=request.user,
        )
        copied += 1
    log_action(action="budget_version_created", entity_type="budgetversion",
               entity_id=new_version.pk,
               diff={"budget": str(budget), "version": new_version.version,
                     "copied_from": source.version, "lines_copied": copied,
                     "reason": "نسخة مراجعة (Revised Budget)"},
               request=request)
    messages.success(
        request,
        f"أُنشئت نسخة المراجعة v{new_version.version} بنسخ {copied} سطرًا من "
        f"v{source.version} — عدّل ما يلزم ثم اعتمد النسخة الجديدة.",
    )
    return redirect("budget:version_detail", pk=new_version.pk)


# ---------------------------------------------------------------- lines
def _require_draft(version, action_label: str):
    if version.is_approved:
        raise PermissionDenied(
            f"النسخة {version} معتمدة ومُقفلة — {action_label} متاح فقط على "
            "المسودات. أنشئ نسخة مراجعة (Revised Budget) ثم عدّل فيها."
        )


@require_permission("budget.edit")
def line_create(request, version_pk):
    version = get_object_or_404(BudgetVersion, pk=version_pk)
    _require_draft(version, "إضافة سطر")
    if request.method == "POST":
        form = forms.BudgetLineForm(request.POST)
        if form.is_valid():
            line = form.save(commit=False)
            line.version = version
            line.created_by = request.user
            line.updated_by = request.user
            line.save()
            log_action(action="entity_created", entity_type="budgetline",
                       entity_id=line.pk, diff=snapshot(line), request=request)
            messages.success(
                request,
                f"أُضيف السطر {line} (السنوية {line.annual_amount}) وسُجِّل في سجل التدقيق.",
            )
            return redirect("budget:version_detail", pk=version.pk)
    else:
        form = forms.BudgetLineForm()
    return render(request, "budget/line_form.html", {
        "form": form, "version": version, "is_create": True,
        "cancel_url": reverse("budget:version_detail", args=[version.pk]),
    })


@require_permission("budget.edit")
def line_edit(request, pk):
    line = get_object_or_404(
        BudgetLine.objects.select_related("version__budget"), pk=pk)
    _require_draft(line.version, "تعديل سطر")
    before = snapshot(line)
    if request.method == "POST":
        form = forms.BudgetLineForm(request.POST, instance=line)
        if form.is_valid():
            line = form.save(commit=False)
            line.updated_by = request.user
            line.save()
            diff = changes_between(line, before)
            if diff:
                log_action(action="entity_updated", entity_type="budgetline",
                           entity_id=line.pk, diff=diff, request=request)
                messages.success(
                    request,
                    f"عُدّل السطر {line} (السنوية الجديدة {line.annual_amount}) "
                    "والفروق مُسجَّلة.",
                )
            return redirect("budget:version_detail", pk=line.version_id)
    else:
        form = forms.BudgetLineForm(instance=line)
    return render(request, "budget/line_form.html", {
        "form": form, "version": line.version, "is_create": False, "line": line,
        "cancel_url": reverse("budget:version_detail", args=[line.version_id]),
    })


@require_permission("budget.edit")
def line_delete(request, pk):
    line = get_object_or_404(
        BudgetLine.objects.select_related("version__budget"), pk=pk)
    _require_draft(line.version, "حذف سطر")
    if request.method != "POST":
        messages.error(request, "الحذف يتطلب تأكيدًا (POST).")
        return redirect("budget:version_detail", pk=line.version_id)
    before = snapshot(line)
    label, version_id = str(line), line.version_id
    line.delete()
    log_action(action="entity_deleted", entity_type="budgetline", entity_id=pk,
               diff={"snapshot": before, "label": label}, request=request)
    messages.success(request, f"حُذف السطر «{label}» وسُجِّل في سجل التدقيق.")
    return redirect("budget:version_detail", pk=version_id)


# ---------------------------------------------------------------- summary
def version_summary(request, pk):
    """Budget-vs-Actual foundation: DB-computed totals for ONE version,
    with an explicit banner stating whether this version is THE analysis
    version. (Actuals/Variance arrive in a later phase.)"""
    version = get_object_or_404(
        BudgetVersion.objects.select_related("budget__fiscal_year"), pk=pk)
    totals = services.version_totals(version)
    context = {
        "version": version, "budget": version.budget, "totals": totals,
        "by_department": services.totals_by(version, "department"),
        "by_account": services.totals_by(version, "account"),
        "by_category": services.totals_by(version, "expense_category"),
        "is_analysis": version.is_analysis_version,
        "back_url": reverse("budget:version_detail", args=[version.pk]),
    }
    return render(request, "budget/summary.html", context)


# ---------------------------------------------------------------- configurable budget templates
def template_list(request):
    q = (request.GET.get("q") or "").strip()
    queryset = BudgetTemplate.objects.select_related("currency").annotate(
        sheet_count=Count("sheets", distinct=True),
        mapped_rows=Count("sheets__rows", filter=Q(sheets__rows__account__isnull=False)),
    )
    if q:
        queryset = queryset.filter(
            Q(name__icontains=q) | Q(source_filename__icontains=q)
        )
    return render(request, "budget/template_list.html", {
        "templates": queryset,
        "q": q,
        "can_edit": has_perm(request.user, "budget.edit"),
    })


@require_permission("budget.edit")
def template_import(request):
    if request.method == "POST":
        form = forms.BudgetTemplateImportForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                template = import_budget_workbook(
                    form.cleaned_data["workbook"],
                    name=form.cleaned_data["name"],
                    description=form.cleaned_data["description"],
                    user=request.user,
                )
            except Exception as exc:
                form.add_error(
                    "workbook",
                    f"تعذر قراءة المصنف. تأكد من سلامة ملف Excel: {exc}",
                )
            else:
                log_action(
                    action="budget_template_imported",
                    entity_type="budgettemplate",
                    entity_id=template.pk,
                    diff={
                        "name": template.name,
                        "source_filename": template.source_filename,
                        "sheet_count": template.sheets.count(),
                    },
                    request=request,
                )
                messages.success(
                    request,
                    f"تم استيراد النموذج «{template.name}» مع {template.sheets.count()} ورقة.",
                )
                return redirect("budget:template_detail", pk=template.pk)
    else:
        form = forms.BudgetTemplateImportForm()
    return render(request, "budget/template_import.html", {"form": form})


def template_detail(request, pk):
    template = get_object_or_404(
        BudgetTemplate.objects.select_related("currency"), pk=pk,
    )
    sheets = template.sheets.annotate(
        populated_rows=Count("rows", distinct=True),
        mapped_rows=Count(
            "rows", filter=Q(rows__analytical_account__isnull=False), distinct=True,
        ),
    )
    return render(request, "budget/template_detail.html", {
        "template": template,
        "sheets": sheets,
        "can_edit": has_perm(request.user, "budget.edit"),
        "refresh_form": forms.BudgetTemplateRefreshForm(),
    })


@require_permission("budget.edit")
@require_POST
def template_refresh(request, pk):
    template = get_object_or_404(BudgetTemplate, pk=pk)
    form = forms.BudgetTemplateRefreshForm(request.POST, request.FILES)
    if not form.is_valid():
        messages.error(request, "تعذر قراءة نسخة Excel الجديدة. تحقق من الملف وحجمه.")
        return redirect("budget:template_detail", pk=template.pk)
    try:
        stats = refresh_budget_workbook(
            template, form.cleaned_data["workbook"], user=request.user,
        )
    except Exception:
        messages.error(request, "تعذر تحديث النموذج من Excel. تحقق من بنية الملف ثم أعد المحاولة.")
        return redirect("budget:template_detail", pk=template.pk)
    log_action(
        action="budget_template_imported", entity_type="budgettemplate",
        entity_id=template.pk, diff={"refresh": True, **stats}, request=request,
    )
    messages.success(
        request,
        f"تم تحديث {stats['sheets']} ورقة و{stats['rows']} صفًا مع الحفاظ على روابط الحسابات للصفوف المطابقة.",
    )
    return redirect("budget:template_detail", pk=template.pk)


@require_permission("budget.edit")
def template_edit(request, pk):
    template = get_object_or_404(BudgetTemplate, pk=pk)
    before = snapshot(template, ["name", "description", "currency", "is_active"])
    if request.method == "POST":
        form = forms.BudgetTemplateForm(request.POST, instance=template)
        if form.is_valid():
            template = form.save(commit=False)
            template.updated_by = request.user
            template.save()
            diff = changes_between(template, before)
            if diff:
                log_action(
                    action="entity_updated", entity_type="budgettemplate",
                    entity_id=template.pk, diff=diff, request=request,
                )
            messages.success(request, "تم تحديث إعدادات نموذج الموازنة.")
            return redirect("budget:template_detail", pk=template.pk)
    else:
        form = forms.BudgetTemplateForm(instance=template)
    return render(request, "budget/template_edit.html", {
        "form": form,
        "template": template,
    })


@require_permission("budget.edit")
@require_POST
def template_delete(request, pk):
    template = get_object_or_404(BudgetTemplate, pk=pk)
    source_filename = template.source_filename
    label = template.name
    template.delete()
    log_action(
        action="entity_deleted", entity_type="budgettemplate", entity_id=pk,
        diff={"name": label, "source_filename": source_filename}, request=request,
    )
    messages.success(request, f"تم حذف نموذج الموازنة «{label}».")
    return redirect("budget:template_list")


def _normalize_template_row_positions(sheet):
    rows = list(sheet.rows.order_by("position", "row_number", "pk"))
    changed = []
    for position, row in enumerate(rows, start=1):
        if row.position != position:
            row.position = position
            changed.append(row)
    if changed:
        BudgetTemplateRow.objects.bulk_update(changed, ["position"])
    return rows


def _sum_previous_template_rows(sheet, target_row, columns):
    rows = _normalize_template_row_positions(sheet)
    target_index = next(
        (index for index, row in enumerate(rows) if row.pk == target_row.pk), None,
    )
    if target_index is None:
        return 0, 0
    boundaries = {
        BudgetTemplateRow.TYPE_TITLE,
        BudgetTemplateRow.TYPE_HEADER,
        BudgetTemplateRow.TYPE_SUBTOTAL,
        BudgetTemplateRow.TYPE_TOTAL,
    }
    source_rows = []
    for row in reversed(rows[:target_index]):
        if row.row_type in boundaries:
            break
        if row.row_type == BudgetTemplateRow.TYPE_DATA:
            source_rows.append(row)
    source_rows.reverse()
    if not source_rows:
        return 0, 0

    sum_columns = [
        column for column in columns
        if column.role == BudgetTemplateColumn.ROLE_ANNUAL
        or column.role.startswith("month_")
    ]
    changed_cells = 0
    for column in sum_columns:
        coordinates = list(BudgetTemplateCell.objects.filter(
            row__in=source_rows, column=column,
        ).order_by("row__position").values_list("coordinate", flat=True))
        if not coordinates:
            continue
        formula = f"=SUM({','.join(coordinates)})"
        cell, _ = BudgetTemplateCell.objects.get_or_create(
            row=target_row,
            column=column,
            defaults={"coordinate": f"{column.column_letter}{target_row.row_number}"},
        )
        cell.raw_value = formula
        cell.formula = formula
        cell.override_value = None
        cell.is_editable = False
        cell.save(update_fields=[
            "raw_value", "formula", "override_value", "is_editable",
        ])
        changed_cells += 1
    return len(source_rows), changed_cells


def template_sheet(request, pk):
    sheet = get_object_or_404(
        BudgetTemplateSheet.objects.select_related("template"), pk=pk,
    )
    columns = list(sheet.columns.all())
    row_queryset = sheet.rows.select_related(
        "account", "main_account", "analytical_account",
    ).prefetch_related(
        "cells__column"
    )
    paginator = Paginator(row_queryset, 60)
    page = paginator.get_page(request.GET.get("page"))
    can_edit = has_perm(request.user, "budget.edit")
    page_url = f"{reverse('budget:template_sheet', args=[sheet.pk])}?page={page.number}"

    action = request.POST.get("action") if request.method == "POST" else ""
    if request.method == "POST" and not can_edit:
        raise PermissionDenied("لا تملك صلاحية تعديل نماذج الموازنة.")

    if action in {
        "add_row", "duplicate_row", "delete_row", "move_up", "move_down",
        "sum_previous",
    }:
        with transaction.atomic():
            _normalize_template_row_positions(sheet)
            if action == "add_row":
                row_number = (sheet.rows.aggregate(value=models.Max("row_number"))["value"] or 0) + 1
                position = (sheet.rows.aggregate(value=models.Max("position"))["value"] or 0) + 1
                new_row = BudgetTemplateRow.objects.create(
                    sheet=sheet, row_number=row_number, position=position,
                    row_type=BudgetTemplateRow.TYPE_DATA,
                    label="صف جديد", is_included=True,
                )
                BudgetTemplateCell.objects.bulk_create([
                    BudgetTemplateCell(
                        row=new_row, column=column,
                        coordinate=f"{column.column_letter}{row_number}",
                        is_editable=True,
                    )
                    for column in columns
                ])
                messages.success(request, "تمت إضافة صف جديد في نهاية الورقة.")
            else:
                source_row = get_object_or_404(
                    BudgetTemplateRow, pk=request.POST.get("row_id"), sheet=sheet,
                )
                if action == "delete_row":
                    source_row.delete()
                    _normalize_template_row_positions(sheet)
                    messages.success(request, "تم حذف الصف من النموذج.")
                elif action in {"move_up", "move_down"}:
                    direction = -1 if action == "move_up" else 1
                    adjacent = sheet.rows.filter(
                        position=source_row.position + direction,
                    ).first()
                    if adjacent:
                        source_position = source_row.position
                        source_row.position = adjacent.position
                        adjacent.position = source_position
                        BudgetTemplateRow.objects.bulk_update(
                            [source_row, adjacent], ["position"],
                        )
                        messages.success(request, "تم تغيير ترتيب الصف.")
                    else:
                        messages.info(request, "الصف موجود بالفعل عند حد الترتيب.")
                elif action == "sum_previous":
                    if source_row.row_type not in {
                        BudgetTemplateRow.TYPE_SUBTOTAL,
                        BudgetTemplateRow.TYPE_TOTAL,
                    }:
                        messages.error(
                            request,
                            "أمر جمع السابق متاح لصف الإجمالي أو الإجمالي الفرعي فقط.",
                        )
                    else:
                        source_count, cell_count = _sum_previous_template_rows(
                            sheet, source_row, columns,
                        )
                        if cell_count:
                            messages.success(
                                request,
                                f"تم ربط الإجمالي تلقائيًا بـ {source_count} صف سابق في {cell_count} عمود مالي.",
                            )
                        else:
                            messages.warning(
                                request,
                                "لم توجد صفوف بيانات سابقة أو أعمدة شهرية/سنوية محددة للجمع.",
                            )
                else:
                    from openpyxl.formula.translate import Translator

                    row_number = (sheet.rows.aggregate(value=models.Max("row_number"))["value"] or 0) + 1
                    BudgetTemplateRow.objects.filter(
                        sheet=sheet, position__gt=source_row.position,
                    ).update(position=models.F("position") + 1)
                    new_row = BudgetTemplateRow.objects.create(
                        sheet=sheet, row_number=row_number,
                        position=source_row.position + 1,
                        row_type=source_row.row_type, label=source_row.label,
                        account=source_row.account,
                        main_account=source_row.main_account,
                        analytical_account=source_row.analytical_account,
                        is_included=source_row.is_included,
                        height=source_row.height,
                    )
                    copied_cells = []
                    for source_cell in source_row.cells.select_related("column"):
                        coordinate = f"{source_cell.column.column_letter}{row_number}"
                        formula = source_cell.formula
                        if formula:
                            try:
                                formula = Translator(
                                    formula, origin=source_cell.coordinate,
                                ).translate_formula(coordinate)
                            except Exception:
                                pass
                        copied_cells.append(BudgetTemplateCell(
                            row=new_row, column=source_cell.column,
                            coordinate=coordinate,
                            raw_value=formula or source_cell.effective_value,
                            formula=formula,
                            data_type=source_cell.data_type,
                            number_format=source_cell.number_format,
                            style_metadata=source_cell.style_metadata,
                            is_editable=source_cell.is_editable,
                        ))
                    BudgetTemplateCell.objects.bulk_create(copied_cells)
                    messages.success(request, "تم نسخ الصف أسفل الصف الأصلي.")
        return redirect(page_url)

    if action == "reset_page":
        BudgetTemplateCell.objects.filter(row__in=page.object_list).update(
            override_value=None,
        )
        messages.success(request, "تمت استعادة القيم الأصلية لخلايا الصفحة الحالية.")
        return redirect(page_url)

    if request.method == "POST":
        sheet_form = forms.BudgetTemplateSheetForm(
            request.POST, instance=sheet, prefix="sheet",
        )
        column_forms = [
            forms.BudgetTemplateColumnRoleForm(
                request.POST, instance=column, prefix=f"column-{column.pk}",
            )
            for column in columns
        ]
        row_forms = [
            forms.BudgetTemplateRowMappingForm(
                request.POST, instance=row, prefix=f"row-{row.pk}"
            )
            for row in page.object_list
        ]
        if (
            sheet_form.is_valid()
            and all(form.is_valid() for form in column_forms)
            and all(form.is_valid() for form in row_forms)
        ):
            with transaction.atomic():
                changed_rows = 0
                changed_columns = 0
                changed_cells = 0
                old_purpose = sheet.purpose
                updated_sheet = sheet_form.save(commit=False)
                updated_sheet.updated_by = request.user
                updated_sheet.save()
                for column_form in column_forms:
                    if column_form.has_changed():
                        column_form.save()
                        changed_columns += 1
                for row_form in row_forms:
                    if row_form.has_changed():
                        mapped_row = row_form.save(commit=False)
                        if mapped_row.row_type == BudgetTemplateRow.TYPE_DATA:
                            mapped_row.account = (
                                mapped_row.analytical_account or mapped_row.main_account
                            )
                        else:
                            mapped_row.account = None
                            mapped_row.main_account = None
                            mapped_row.analytical_account = None
                            mapped_row.is_included = False
                        mapped_row.save()
                        changed_rows += 1
                allowed_cells = {
                    cell.pk: cell
                    for row in page.object_list
                    for cell in row.cells.all()
                    if not cell.formula
                }
                for key, value in request.POST.items():
                    if not key.startswith("cell-") or not key[5:].isdigit():
                        continue
                    cell = allowed_cells.get(int(key[5:]))
                    if cell is None:
                        continue
                    override = None if value == cell.raw_value else value
                    if cell.override_value != override:
                        cell.override_value = override
                        cell.save(update_fields=["override_value"])
                        changed_cells += 1
            log_action(
                action="budget_template_mapping_updated",
                entity_type="budgettemplatesheet",
                entity_id=sheet.pk,
                diff={
                    "purpose": {"before": old_purpose, "after": sheet.purpose},
                    "changed_rows": changed_rows,
                    "changed_columns": changed_columns,
                    "changed_cells": changed_cells,
                    "page": page.number,
                },
                request=request,
            )
            messages.success(
                request,
                f"تم حفظ الورقة: {changed_cells} خلية، و{changed_rows} ربط صف، و{changed_columns} وظيفة عمود.",
            )
            return redirect(page_url)
    else:
        sheet_form = forms.BudgetTemplateSheetForm(instance=sheet, prefix="sheet")
        column_forms = [
            forms.BudgetTemplateColumnRoleForm(
                instance=column, prefix=f"column-{column.pk}",
            )
            for column in columns
        ]
        row_forms = [
            forms.BudgetTemplateRowMappingForm(instance=row, prefix=f"row-{row.pk}")
            for row in page.object_list
        ]

    from openpyxl.utils.cell import range_boundaries

    merge_anchors = {}
    merged_covered = set()
    for merged_range in sheet.merged_ranges:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(merged_range)
        except ValueError:
            continue
        anchor = f"{columns[min_col - 1].column_letter}{min_row}" if min_col <= len(columns) else ""
        merge_anchors[anchor] = {
            "colspan": max_col - min_col + 1,
            "rowspan": max_row - min_row + 1,
        }
        for row_number in range(min_row, max_row + 1):
            for column_number in range(min_col, max_col + 1):
                coordinate = f"{columns[column_number - 1].column_letter}{row_number}" if column_number <= len(columns) else ""
                if coordinate and coordinate != anchor:
                    merged_covered.add(coordinate)

    runtime = template_runtime.SheetRuntime(sheet)

    def cell_style(cell):
        if cell is None:
            return ""
        metadata = cell.style_metadata or {}
        styles = []
        if metadata.get("bold"):
            styles.append("font-weight:700")
        if metadata.get("italic"):
            styles.append("font-style:italic")
        font_color = str(metadata.get("font_color") or "")[-6:]
        fill_color = str(metadata.get("fill_color") or "")[-6:]
        if len(font_color) == 6 and font_color != "000000":
            styles.append(f"color:#{font_color}")
        if len(fill_color) == 6 and fill_color not in {"000000", "FFFFFF"}:
            styles.append(f"background-color:#{fill_color}")
        if metadata.get("horizontal"):
            styles.append(f"text-align:{metadata['horizontal']}")
        if metadata.get("wrap_text"):
            styles.append("white-space:normal")
        return ";".join(styles)

    rows = []
    for display_number, (row, row_form) in enumerate(
        zip(page.object_list, row_forms), start=page.start_index(),
    ):
        cells_by_column = {cell.column_id: cell for cell in row.cells.all()}
        rendered_cells = []
        for column in columns:
            cell = cells_by_column.get(column.pk)
            coordinate = cell.coordinate if cell else f"{column.column_letter}{row.row_number}"
            rendered_cells.append({
                "cell": cell,
                "coordinate": coordinate,
                "skip": coordinate in merged_covered,
                "merge": merge_anchors.get(coordinate, {}),
                "value": (
                    request.POST.get(f"cell-{cell.pk}", runtime.display_value(cell))
                    if cell else ""
                ),
                "style": cell_style(cell),
            })
        rows.append({
            "row": row,
            "form": row_form,
            "cells": rendered_cells,
            "display_number": display_number,
        })
    return render(request, "budget/template_sheet.html", {
        "sheet": sheet,
        "template": sheet.template,
        "columns": columns,
        "column_forms": column_forms,
        "rows": rows,
        "page_obj": page,
        "sheet_form": sheet_form,
        "can_edit": can_edit,
    })


def template_sheet_monthly_output(request, pk):
    sheet = get_object_or_404(BudgetTemplateSheet.objects.select_related("template"), pk=pk)
    return render(request, "budget/template_runtime_output.html", {
        "sheet": sheet,
        "template": sheet.template,
        "rows": template_runtime.template_monthly_output(sheet),
        "title": "التوزيع الشهري للنموذج المرن",
        "summary": False,
    })


def template_sheet_summary_output(request, pk):
    sheet = get_object_or_404(BudgetTemplateSheet.objects.select_related("template"), pk=pk)
    return render(request, "budget/template_runtime_output.html", {
        "sheet": sheet,
        "template": sheet.template,
        "rows": template_runtime.template_summary_output(sheet),
        "title": "إجماليات الحسابات الرئيسية للنموذج المرن",
        "summary": True,
    })


def _ensure_plan_editable(plan: BudgetPlan):
    if plan.is_locked:
        raise PermissionDenied("النموذج المعتمد أو المؤرشف مقفل؛ أنشئ إصدارًا جديدًا للتعديل.")


def plan_list(request):
    plans = BudgetPlan.objects.select_related("fiscal_year", "currency").annotate(
        section_count=Count("sections"),
    )
    return render(request, "budget/plan_list.html", {
        "plans": plans,
        "can_edit": has_perm(request.user, "budget.edit"),
    })


@require_permission("budget.edit")
def plan_create(request):
    form = forms.BudgetPlanForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        plan = form.save(commit=False)
        plan.created_by = request.user
        plan.updated_by = request.user
        plan.save()
        log_action(
            action="budget_plan_created", entity_type="budgetplan",
            entity_id=plan.pk, diff=snapshot(plan), request=request,
        )
        messages.success(request, "تم إنشاء نموذج التخطيط التشغيلي.")
        return redirect("budget:plan_detail", pk=plan.pk)
    return render(request, "budget/plan_form.html", {"form": form, "is_create": True})


def plan_detail(request, pk):
    plan = get_object_or_404(
        BudgetPlan.objects.select_related("fiscal_year", "currency"), pk=pk,
    )
    return render(request, "budget/plan_detail.html", {
        "plan": plan,
        "sections": services.detailed_plan_output(plan),
        "can_edit": has_perm(request.user, "budget.edit") and not plan.is_locked,
    })


@require_permission("budget.edit")
def plan_edit(request, pk):
    plan = get_object_or_404(BudgetPlan, pk=pk)
    _ensure_plan_editable(plan)
    before = snapshot(plan)
    form = forms.BudgetPlanForm(request.POST or None, instance=plan)
    if request.method == "POST" and form.is_valid():
        plan = form.save(commit=False)
        plan.updated_by = request.user
        plan.save()
        log_action(
            action="budget_plan_updated", entity_type="budgetplan",
            entity_id=plan.pk, diff=changes_between(plan, before), request=request,
        )
        messages.success(request, "تم تحديث نموذج التخطيط.")
        return redirect("budget:plan_detail", pk=plan.pk)
    return render(request, "budget/plan_form.html", {"form": form, "plan": plan})


@require_permission("budget.edit")
@require_POST
def plan_delete(request, pk):
    plan = get_object_or_404(BudgetPlan, pk=pk)
    _ensure_plan_editable(plan)
    plan_id, label = plan.pk, str(plan)
    plan.delete()
    log_action(
        action="budget_plan_deleted", entity_type="budgetplan",
        entity_id=plan_id, diff={"label": label}, request=request,
    )
    messages.success(request, "تم حذف نموذج التخطيط.")
    return redirect("budget:plan_list")


@require_permission("budget.edit")
def plan_section_create(request, plan_pk):
    plan = get_object_or_404(BudgetPlan, pk=plan_pk)
    _ensure_plan_editable(plan)
    form = forms.BudgetPlanSectionForm(request.POST or None, plan=plan)
    if request.method == "POST" and form.is_valid():
        section = form.save(commit=False)
        section.plan = plan
        section.created_by = request.user
        section.updated_by = request.user
        section.save()
        log_action(
            action="budget_plan_section_created", entity_type="budgetplansection",
            entity_id=section.pk, diff=snapshot(section), request=request,
        )
        return redirect("budget:plan_detail", pk=plan.pk)
    return render(request, "budget/plan_section_form.html", {"form": form, "plan": plan})


@require_permission("budget.edit")
def plan_section_edit(request, pk):
    section = get_object_or_404(
        BudgetPlanSection.objects.select_related("plan", "department"), pk=pk,
    )
    _ensure_plan_editable(section.plan)
    form = forms.BudgetPlanSectionForm(
        request.POST or None, instance=section, plan=section.plan,
    )
    if request.method == "POST" and form.is_valid():
        section = form.save(commit=False)
        section.updated_by = request.user
        section.save()
        log_action(
            action="budget_plan_section_updated", entity_type="budgetplansection",
            entity_id=section.pk, diff=snapshot(section), request=request,
        )
        return redirect("budget:plan_detail", pk=section.plan_id)
    return render(request, "budget/plan_section_form.html", {"form": form, "plan": section.plan, "section": section})


@require_permission("budget.edit")
@require_POST
def plan_section_delete(request, pk):
    section = get_object_or_404(BudgetPlanSection.objects.select_related("plan"), pk=pk)
    _ensure_plan_editable(section.plan)
    plan_id = section.plan_id
    section_id, label = section.pk, str(section)
    section.delete()
    log_action(
        action="budget_plan_section_deleted", entity_type="budgetplansection",
        entity_id=section_id, diff={"label": label}, request=request,
    )
    messages.success(request, "تم حذف القسم وبنوده.")
    return redirect("budget:plan_detail", pk=plan_id)


def _line_form_response(request, section, line=None):
    initial = None
    if request.method == "GET" and line is None:
        initial = {}
        parent_id = request.GET.get("parent")
        if parent_id and section.lines.filter(pk=parent_id).exists():
            initial["parent"] = parent_id
    line_instance = line or BudgetPlanLine(section=section)
    form = forms.BudgetPlanLineForm(
        request.POST or None, instance=line_instance, section=section, initial=initial,
    )
    salary_formset_data = None
    if request.method == "POST" and (
        request.POST.get("input_mode") == BudgetPlanLine.INPUT_SALARY_COMPONENTS
        or "salary-TOTAL_FORMS" in request.POST
    ):
        salary_formset_data = request.POST
    salary_formset = forms.BudgetPlanSalaryComponentFormSet(
        salary_formset_data, instance=line_instance, prefix="salary",
    )
    salary_formset_is_valid = not salary_formset.is_bound or salary_formset.is_valid()
    if request.method == "POST" and form.is_valid() and salary_formset_is_valid:
        try:
            with transaction.atomic():
                saved_line = form.save(commit=False)
                saved_line.section = section
                if not saved_line.pk:
                    saved_line.created_by = request.user
                saved_line.updated_by = request.user
                saved_line.save()
                form.save_m2m()
                if salary_formset.is_bound:
                    salary_formset.instance = saved_line
                    salary_components = salary_formset.save(commit=False)
                    for deleted_component in salary_formset.deleted_objects:
                        deleted_component.delete()
                    for component in salary_components:
                        if not component.pk:
                            component.created_by = request.user
                        component.updated_by = request.user
                        component.save()
                services.sync_plan_line_amounts(saved_line, form.monthly_values())
        except ValidationError as exc:
            for field, errors in exc.message_dict.items():
                target = field if field in form.fields else None
                for error in errors:
                    form.add_error(target, error)
        else:
            log_action(
                action="budget_plan_line_updated" if line else "budget_plan_line_created",
                entity_type="budgetplanline", entity_id=saved_line.pk,
                diff=snapshot(saved_line), request=request,
            )
            messages.success(request, "تم حفظ بند الموازنة وإعادة احتساب توزيعه.")
            return redirect("budget:plan_detail", pk=section.plan_id)
    return render(request, "budget/plan_line_form.html", {
        "form": form, "section": section, "plan": section.plan, "line": line,
        "salary_formset": salary_formset,
        "non_month_fields": [
            field for field in form if field.name not in MONTH_FIELDS
        ],
        "month_fields": [form[f"m{month:02d}"] for month in range(1, 13)],
    })


@require_permission("budget.edit")
def plan_line_create(request, section_pk):
    section = get_object_or_404(BudgetPlanSection.objects.select_related("plan"), pk=section_pk)
    _ensure_plan_editable(section.plan)
    return _line_form_response(request, section)


@require_permission("budget.edit")
def plan_line_edit(request, pk):
    line = get_object_or_404(BudgetPlanLine.objects.select_related("section__plan"), pk=pk)
    _ensure_plan_editable(line.plan)
    return _line_form_response(request, line.section, line)


@require_permission("budget.edit")
def plan_section_bulk_lines(request, pk):
    section = get_object_or_404(
        BudgetPlanSection.objects.select_related("plan", "department"), pk=pk,
    )
    _ensure_plan_editable(section.plan)
    imported_rows = []
    import_warnings = []
    imported_sheet = ""
    action = request.POST.get("action") if request.method == "POST" else ""

    if action == "preview_excel":
        upload = request.FILES.get("excel_file")
        if upload is None:
            messages.error(request, "اختر ملف Excel أولًا.")
        elif not upload.name.lower().endswith(".xlsx"):
            messages.error(request, "الصيغة المدعومة في هذه المرحلة هي xlsx فقط.")
        elif upload.size > 10 * 1024 * 1024:
            messages.error(request, "حجم الملف يتجاوز الحد المسموح (10 ميجابايت).")
        else:
            try:
                imported_rows, import_warnings, imported_sheet = parse_plan_lines_workbook(
                    upload, section,
                )
            except PlanLineImportError as exc:
                messages.error(request, str(exc))
            else:
                messages.success(
                    request,
                    f"تمت قراءة {len(imported_rows)} بندًا من ورقة «{imported_sheet}». راجعها ثم اضغط حفظ الكل.",
                )

    extra = len(imported_rows) if imported_rows else 5
    FormSet = forms.budget_plan_bulk_line_formset(extra=extra)
    queryset = section.lines.select_related("main_account", "analytical_account").order_by(
        "position", "pk",
    )
    if action == "save":
        formset = FormSet(
            request.POST, queryset=queryset, prefix="lines",
            form_kwargs={"section": section},
        )
        if formset.is_valid():
            saved_count = 0
            deleted_count = 0
            try:
                with transaction.atomic():
                    section.lines.update(position=models.F("position") + 1000000)
                    for form in formset.forms:
                        if not form.cleaned_data:
                            continue
                        instance = form.instance
                        if form.cleaned_data.get("DELETE"):
                            if instance.pk:
                                instance.delete()
                                deleted_count += 1
                            continue
                        instance = form.save(commit=False)
                        instance.section = section
                        instance.department = section.department
                        if not instance.pk:
                            instance.created_by = request.user
                        instance.updated_by = request.user
                        instance.full_clean()
                        instance.save()
                        services.sync_plan_line_amounts(instance)
                        saved_count += 1
            except (ValidationError, ProtectedError) as exc:
                messages.error(request, f"تعذر حفظ البنود: {exc}")
            else:
                log_action(
                    action="budget_plan_line_updated", entity_type="budgetplansection",
                    entity_id=section.pk,
                    diff={"bulk_saved": saved_count, "bulk_deleted": deleted_count},
                    request=request,
                )
                messages.success(
                    request,
                    f"تم حفظ {saved_count} بندًا" + (
                        f" وحذف {deleted_count} بندًا." if deleted_count else "."
                    ),
                )
                return redirect(f"{reverse('budget:plan_detail', args=[section.plan_id])}#section-{section.pk}")
    else:
        formset = FormSet(
            queryset=queryset, initial=imported_rows, prefix="lines",
            form_kwargs={"section": section},
        )

    return render(request, "budget/plan_section_bulk_lines.html", {
        "plan": section.plan,
        "section": section,
        "formset": formset,
        "import_warnings": import_warnings,
        "imported_sheet": imported_sheet,
    })


@require_permission("budget.edit")
@require_GET
def plan_section_import_template(request, pk):
    import re

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    section = get_object_or_404(BudgetPlanSection.objects.select_related("plan"), pk=pk)
    _ensure_plan_editable(section.plan)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = re.sub(r"[\\/*?:\[\]]", "-", section.name)[:31] or "بنود القسم"
    headers = [
        "الترتيب", "نوع الصف", "اسم البند", "الحساب الرئيسي",
        "الحساب التحليلي", "طريقة الإدخال", "طريقة التوزيع",
        "المبلغ السنوي", "أساس التقدير", "مضمن",
    ]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.rightToLeft = True
    worksheet.column_dimensions["C"].width = 34
    worksheet.column_dimensions["D"].width = 24
    worksheet.column_dimensions["E"].width = 28
    worksheet.column_dimensions["I"].width = 36
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="budget-section-lines-template.xlsx"'
    workbook.save(response)
    workbook.close()
    return response


@require_permission("budget.edit")
@require_POST
def plan_line_delete(request, pk):
    line = get_object_or_404(BudgetPlanLine.objects.select_related("section__plan"), pk=pk)
    _ensure_plan_editable(line.plan)
    plan_id = line.plan.pk
    line_id, label = line.pk, str(line)
    try:
        line.delete()
    except ProtectedError:
        messages.error(request, "لا يمكن حذف البند لوجود بنود فرعية مرتبطة به.")
        return redirect("budget:plan_detail", pk=plan_id)
    log_action(
        action="budget_plan_line_deleted", entity_type="budgetplanline",
        entity_id=line_id, diff={"label": label}, request=request,
    )
    messages.success(request, "تم حذف البند.")
    return redirect("budget:plan_detail", pk=plan_id)


@require_permission("budget.edit")
@require_POST
def plan_line_copy(request, pk):
    source = get_object_or_404(BudgetPlanLine.objects.select_related("section__plan"), pk=pk)
    _ensure_plan_editable(source.plan)
    with transaction.atomic():
        last_position = source.section.lines.aggregate(value=models.Max("position"))["value"] or 0
        monthly = {row.month: row.amount for row in source.period_amounts.all()}
        aggregate_section_ids = list(source.aggregate_sections.values_list("pk", flat=True))
        salary_component_values = list(source.salary_components.values(
            "name", "component_type", "calculation_method", "amount",
            "percentage_rate", "start_month", "periods_count", "payment_month",
            "is_percentage_base", "position", "is_active",
        ))
        source.pk = None
        source.position = last_position + 1
        source.display_name = f"{source.display_name} - نسخة"
        source.created_by = request.user
        source.updated_by = request.user
        source.save()
        source.aggregate_sections.set(aggregate_section_ids)
        for component_values in salary_component_values:
            source.salary_components.create(
                **component_values,
                created_by=request.user,
                updated_by=request.user,
            )
        services.sync_plan_line_amounts(source, monthly)
        log_action(
            action="budget_plan_line_copied", entity_type="budgetplanline",
            entity_id=source.pk, diff={"source_id": pk}, request=request,
        )
    messages.success(request, "تم نسخ البند.")
    return redirect("budget:plan_detail", pk=source.plan.pk)


@require_permission("budget.edit")
@require_POST
def plan_line_move(request, pk, direction):
    line = get_object_or_404(BudgetPlanLine.objects.select_related("section__plan"), pk=pk)
    _ensure_plan_editable(line.plan)
    if direction not in {"up", "down"}:
        raise Http404
    siblings = line.section.lines.filter(parent_id=line.parent_id).exclude(pk=line.pk)
    if direction == "up":
        neighbor = siblings.filter(position__lt=line.position).order_by("-position", "-pk").first()
    else:
        neighbor = siblings.filter(position__gt=line.position).order_by("position", "pk").first()
    if neighbor:
        with transaction.atomic():
            old_position = line.position
            temporary_position = (
                line.section.lines.aggregate(value=models.Max("position"))["value"] or 0
            ) + 1
            line.position = temporary_position
            line.updated_by = request.user
            line.save(update_fields=["position", "updated_by", "updated_at"])
            line.position = neighbor.position
            neighbor.position = old_position
            neighbor.updated_by = request.user
            neighbor.save(update_fields=["position", "updated_by", "updated_at"])
            line.save(update_fields=["position", "updated_by", "updated_at"])
        log_action(
            action="budget_plan_line_moved", entity_type="budgetplanline",
            entity_id=line.pk, diff={"direction": direction}, request=request,
        )
        messages.success(request, "تم تغيير ترتيب البند.")
    return redirect("budget:plan_detail", pk=line.plan.pk)


@require_permission("budget.edit")
@require_POST
def plan_line_toggle_included(request, pk):
    line = get_object_or_404(BudgetPlanLine.objects.select_related("section__plan"), pk=pk)
    _ensure_plan_editable(line.plan)
    line.is_included = not line.is_included
    line.updated_by = request.user
    line.save(update_fields=["is_included", "updated_by", "updated_at"])
    log_action(
        action="budget_plan_line_inclusion_toggled", entity_type="budgetplanline",
        entity_id=line.pk, diff={"is_included": line.is_included}, request=request,
    )
    messages.success(
        request,
        "تم تضمين البند في الإجماليات." if line.is_included else "تم استبعاد البند من الإجماليات.",
    )
    return redirect("budget:plan_detail", pk=line.plan.pk)


def plan_monthly_output(request, pk):
    plan = get_object_or_404(BudgetPlan, pk=pk)
    return render(request, "budget/plan_output.html", {
        "plan": plan, "rows": services.monthly_plan_output(plan),
        "title": "الموازنة الشهرية التفصيلية", "show_analytical": True,
    })


def plan_summary_output(request, pk):
    plan = get_object_or_404(BudgetPlan, pk=pk)
    return render(request, "budget/plan_output.html", {
        "plan": plan, "rows": services.summary_plan_output(plan),
        "title": "إجماليات الموازنة الشهرية", "show_analytical": False,
    })


@require_GET
def plan_analytical_accounts(request):
    from apps.reference.models import Account
    main_id = request.GET.get("main_account")
    rows = Account.objects.none()
    if main_id and main_id.isdigit():
        rows = Account.objects.analytical_for(int(main_id)).filter(
            account_type="expense",
        ).order_by("code")
    return JsonResponse({"results": [
        {"id": row.pk, "text": str(row)} for row in rows
    ]})


@require_GET
def plan_department_employees(request):
    from apps.reference.models import Employee
    department_id = request.GET.get("department")
    rows = Employee.objects.none()
    if department_id and department_id.isdigit():
        rows = Employee.objects.filter(
            department_id=int(department_id), is_active=True,
        ).order_by("code")
    return JsonResponse({"results": [
        {"id": row.pk, "text": str(row)} for row in rows
    ]})
