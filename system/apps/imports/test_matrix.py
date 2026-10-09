from __future__ import annotations

from io import BytesIO

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.budget.models import Budget, BudgetLine, BudgetVersion
from apps.expenses.models import Expense
from apps.imports import services
from apps.imports.matrix import template_xlsx
from apps.imports.models import ImportJob
from apps.imports.tests import BaseFixture, TEMP_MEDIA
from apps.reference.models import Account
from apps.reports.builders import budget_for, build_budget_vs_actual, build_variance, filtered_budget_lines


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class AccountMatrixTests(BaseFixture):
    def setUp(self):
        self.budget = Budget.objects.create(fiscal_year=self.fy, name="موازنة الاختبار")
        self.version = BudgetVersion.objects.create(budget=self.budget, version=1)
        self.analytical = Account.objects.create(
            code="510201", name="تفصيل المستلزمات", account_type="expense",
            level=6, ledger_type="sub", parent=self.acc,
        )

    def workbook(self, entries):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["اسم الحساب", *[f"الشهر {index}" for index in range(1, 13)], "الإجمالي"])
        for name, months in entries:
            sheet.append([name, *months, sum(months)])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def job(self, target, entries):
        job = self.extract("matrix.xlsx", self.workbook(entries), fmt="xlsx")
        job.target = target
        job.column_map = services.suggest_mapping(job.headers, target)
        job.column_map["_context"] = {
            "department_id": self.dept.pk,
            "version_id": self.version.pk,
            "period_id": self.p_jan.pk,
        }
        job.save(update_fields=["target", "column_map"])
        return job

    def test_template_has_only_fourteen_columns(self):
        sheet = load_workbook(BytesIO(template_xlsx()), read_only=True).active
        self.assertEqual(sheet.max_column, 14)
        self.assertEqual(sheet["A1"].value, "اسم الحساب")
        self.assertEqual(sheet["N1"].value, "الإجمالي")

    def test_budget_main_and_analytical_are_not_double_counted(self):
        months = [100] + [0] * 11
        job = self.job(ImportJob.TARGET_BUDGET_MATRIX, [
            (self.acc.name, months), (self.analytical.name, months),
        ])
        summary = services.validate_job(job, self.finance)
        self.assertEqual(summary["valid"], 2)
        self.assertEqual(services.run_import(job, self.finance), 1)
        line = BudgetLine.objects.get(version=self.version)
        self.assertEqual(line.analytical_account, self.analytical)
        self.assertEqual(line.annual_amount, 100)

        BudgetLine.objects.create(version=self.version, department=self.dept,
            account=self.acc, m01=100, annual_amount=100)
        self.version.status = "approved"
        self.version.save(update_fields=["status"])
        filters = {"department": None, "account": None, "category": None}
        self.assertEqual(budget_for(filtered_budget_lines(filters, self.fy), [1]), 100)

    def test_monthly_variance_sums_analytical_children(self):
        another = Account.objects.create(
            code="510202", name="تفصيل آخر", account_type="expense",
            level=6, ledger_type="sub", parent=self.acc,
        )
        job = self.job(ImportJob.TARGET_BUDGET_MATRIX, [
            (self.acc.name, [100] + [0] * 11),
            (self.analytical.name, [40] + [0] * 11),
            (another.name, [60] + [0] * 11),
        ])
        self.assertEqual(services.validate_job(job, self.finance)["errors"], 0)
        self.assertEqual(services.run_import(job, self.finance), 2)
        self.version.status = BudgetVersion.STATUS_APPROVED
        self.version.save(update_fields=["status"])
        filters = {"department": self.dept.pk, "account": None, "category": None}
        _, rows, _ = build_variance(
            filters, ([1], self.p_jan.start_date, self.p_jan.end_date, self.fy))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["budget"]["d"], "100.00")

    def test_report_uses_latest_approved_budget_version(self):
        self.version.status = BudgetVersion.STATUS_APPROVED
        self.version.save(update_fields=["status"])
        BudgetLine.objects.create(version=self.version, department=self.dept,
                                  account=self.acc, m01=40, annual_amount=40)
        latest = BudgetVersion.objects.create(budget=self.budget, version=2,
                                              status=BudgetVersion.STATUS_APPROVED)
        BudgetLine.objects.create(version=latest, department=self.dept,
                                  account=self.acc, m01=60, annual_amount=60)
        filters = {"department": self.dept.pk, "account": None, "category": None}
        self.assertEqual(budget_for(filtered_budget_lines(filters, self.fy), [1]), 60)

    def test_main_only_budget_and_monthly_actual_import(self):
        job = self.job(ImportJob.TARGET_BUDGET_MATRIX, [
            (self.acc.name, [70] * 12),
        ])
        services.validate_job(job, self.finance)
        self.assertEqual(services.run_import(job, self.finance), 1)
        self.assertEqual(BudgetLine.objects.get().annual_amount, 840)

        actual = self.job(ImportJob.TARGET_ACTUAL_MATRIX, [
            (self.acc.name, [35] + [0] * 11),
        ])
        services.validate_job(actual, self.finance)
        self.assertEqual(services.run_import(actual, self.finance), 1)
        expense = Expense.objects.get()
        self.assertEqual(expense.amount, 35)
        self.assertEqual(expense.period, self.p_jan)
        self.version.status = "approved"
        self.version.save(update_fields=["status"])
        filters = {"department": self.dept.pk, "account": None, "category": None}
        _, _, summary = build_budget_vs_actual(
            filters, ([1], self.p_jan.start_date, self.p_jan.end_date, self.fy))
        self.assertEqual(dict((item["label"], item["value"]) for item in summary)["Actual Expenses"], "35.00")
        with self.assertRaises(ValidationError):
            services.run_import(actual, self.finance)

    def test_main_only_layout_rejects_analytical_row(self):
        job = self.job(ImportJob.TARGET_BUDGET_MATRIX, [
            (self.analytical.name, [15] + [0] * 11),
        ])
        job.column_map["_context"]["layout"] = "main_only"
        job.save(update_fields=["column_map"])
        self.assertEqual(services.validate_job(job, self.finance)["invalid"], 1)
        with self.assertRaises(ValidationError):
            services.run_import(job, self.finance)

    def test_confirmation_rechecks_new_parent_without_partial_import(self):
        job = self.job(ImportJob.TARGET_BUDGET_MATRIX, [
            (self.analytical.name, [25] + [0] * 11),
        ])
        self.assertEqual(services.validate_job(job, self.finance)["errors"], 0)
        BudgetLine.objects.create(version=self.version, department=self.dept,
                                  account=self.acc, m01=25, annual_amount=25)
        with self.assertRaises(ValidationError):
            services.run_import(job, self.finance)
        self.assertEqual(BudgetLine.objects.count(), 1)

    def test_rejects_wrong_month_and_mismatched_parent(self):
        actual = self.job(ImportJob.TARGET_ACTUAL_MATRIX, [
            (self.acc.name, [0, 30] + [0] * 10),
        ])
        self.assertEqual(services.validate_job(actual, self.finance)["invalid"], 1)
        with self.assertRaises(ValidationError):
            services.run_import(actual, self.finance)

        budget = self.job(ImportJob.TARGET_BUDGET_MATRIX, [
            (self.acc.name, [100] * 12),
            (self.analytical.name, [90] * 12),
        ])
        self.assertGreater(services.validate_job(budget, self.finance)["errors"], 0)
        with self.assertRaises(ValidationError):
            services.run_import(budget, self.finance)

    def test_upload_and_manual_account_mapping_without_extra_file_columns(self):
        client = self.client
        client.force_login(self.finance)
        response = client.post(reverse("imports:new"), {
            "target": ImportJob.TARGET_BUDGET_MATRIX,
            "department": self.dept.pk,
            "budget_version": self.version.pk,
            "matrix_layout": "main_only",
            "source_file": SimpleUploadedFile("accounts.xlsx", self.workbook([
                ("اسم محاسبي قديم", [50] + [0] * 11),
            ])),
        })
        self.assertEqual(response.status_code, 302)
        job = ImportJob.objects.get(target=ImportJob.TARGET_BUDGET_MATRIX)
        self.assertEqual(job.column_map["_context"]["version_id"], self.version.pk)
        detail = client.get(reverse("imports:detail", args=[job.pk]))
        self.assertContains(detail, "اسم محاسبي قديم")
        self.assertContains(detail, "map-account_row_0")
        self.assertContains(detail, 'type="hidden" name="map-m01"')
        self.assertNotContains(detail, 'name="map-fiscal_year"')
        response = client.post(reverse("imports:validate", args=[job.pk]), {
            **{f"map-{key}": header for key, header in job.column_map.items() if key != "_context"},
            "map-account_row_0": str(self.acc.pk),
        })
        self.assertEqual(response.status_code, 302)
        job.refresh_from_db()
        self.assertEqual(job.summary["valid"], 1)
        self.assertEqual(job.column_map["_accounts"]["اسم محاسبي قديم"], str(self.acc.pk))
        response = client.post(reverse("imports:confirm", args=[job.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(BudgetLine.objects.get().annual_amount, 50)
