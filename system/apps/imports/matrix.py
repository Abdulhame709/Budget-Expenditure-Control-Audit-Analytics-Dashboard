"""Import account/month/total workbooks without requiring metadata columns."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction

from apps.budget.models import BudgetLine, BudgetVersion, MONTH_FIELDS
from apps.expenses.models import Expense
from apps.expenses.services import next_expense_number
from apps.governance.services import log_action
from apps.imports.models import ImportJob
from apps.reference.models import Account, Department, MonthlyPeriod
from apps.reference.services import require_period_postable

FIELDS = [("account", "اسم الحساب"),
          *[(month, f"الشهر {index}") for index, month in enumerate(MONTH_FIELDS, 1)],
          ("total", "الإجمالي")]


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "حسابات وشهور"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "B2"
    sheet.append([label for _, label in FIELDS])
    sheet.column_dimensions["A"].width = 38
    for column in sheet.columns:
        column[0].font = Font(bold=True, color="FFFFFF")
        column[0].fill = PatternFill("solid", fgColor="16395B")
    for index in range(2, 15):
        sheet.column_dimensions[get_column_letter(index)].width = 18
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _amount(value, label):
    text = str(value or "").strip().replace(",", "")
    try:
        amount = Decimal(text or "0")
    except (InvalidOperation, ValueError) as exc:
        raise ValidationError(f"{label}: مبلغ غير صالح.") from exc
    if not amount.is_finite() or amount < 0 or amount > Decimal("999999999999.99") or amount.as_tuple().exponent < -2:
        raise ValidationError(f"{label}: مبلغ غير صالح (موجب وبحد أقصى منزلتان عشريتان).")
    return amount


def _account_name(value):
    return " ".join(str(value or "").split()).casefold().replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")


def unmapped_accounts(job):
    header = (job.column_map or {}).get("account") or (job.headers[0] if job.headers else "")
    selected = (job.column_map or {}).get("_accounts", {})
    accounts = Account.objects.filter(is_active=True, account_type="expense", level__in=[5, 6])
    counts = {}
    for account in accounts:
        counts[_account_name(account.name)] = counts.get(_account_name(account.name), 0) + 1
        counts[account.code.casefold()] = counts.get(account.code.casefold(), 0) + 1
    names = {}
    for row in job.rows:
        name = str(row["source"].get(header) or "").strip()
        normalized = _account_name(name)
        if normalized and (counts.get(normalized, 0) != 1 or normalized in selected):
            names.setdefault(normalized, name)
    return names


def _context(job):
    context = (job.column_map or {}).get("_context", {})
    department = Department.objects.filter(pk=context.get("department_id"), is_active=True).first()
    if not department:
        raise ValidationError("اختر إدارة نشطة للملف.")
    if job.target == ImportJob.TARGET_BUDGET_MATRIX:
        scope = BudgetVersion.objects.select_related("budget").filter(pk=context.get("version_id"), status="draft").first()
        if not scope:
            raise ValidationError("نسخة الموازنة غير موجودة أو أصبحت معتمدة.")
    else:
        scope = MonthlyPeriod.objects.select_related("fiscal_year").filter(pk=context.get("period_id"), is_active=True).first()
        if not scope:
            raise ValidationError("الفترة الشهرية غير موجودة أو غير نشطة.")
    return department, scope


def _reference(period, department, account_id):
    return f"MATRIX-{period.pk}-{department.pk}-{account_id}"


@transaction.atomic
def validate_matrix(job, user, request=None):
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status not in (ImportJob.STATUS_EXTRACTED, ImportJob.STATUS_VALIDATED):
        raise ValidationError("لا يمكن التحقق من هذه العملية.")
    if len(job.headers) != 14 or len(set(job.headers)) != 14 or not job.rows:
        raise ValidationError("يتطلب النموذج 14 عمودًا مختلفًا: اسم الحساب، 12 شهرًا، الإجمالي؛ مع صف بيانات واحد على الأقل.")
    mapping = {key: job.column_map.get(key) for key, _ in FIELDS}
    if set(mapping.values()) != set(job.headers) or len(set(mapping.values())) != 14:
        raise ValidationError("طابق اسم الحساب ثم الشهور 1–12 والإجمالي مع أعمدة الملف دون تكرار.")
    department, scope = _context(job)
    is_budget = job.target == ImportJob.TARGET_BUDGET_MATRIX
    if not is_budget:
        require_period_postable(user, scope, purpose="التحقق من ملف المصروفات المجمعة", request=request)
    accounts = Account.objects.filter(is_active=True, account_type="expense", level__in=[5, 6]).select_related("parent")
    by_name = {}
    by_id = {}
    for account in accounts:
        by_id[account.pk] = account
        by_name.setdefault(_account_name(account.name), []).append(account)
        by_name.setdefault(account.code.casefold(), []).append(account)
    seen = set()
    for row in job.rows:
        row.update(mapped={}, ready={}, errors=[], warnings=[], duplicate=False, valid=False)
        mapped = {key: row["source"].get(header, "") for key, header in mapping.items()}
        row["mapped"] = mapped
        try:
            override = (job.column_map.get("_accounts") or {}).get(_account_name(mapped["account"]))
            matches = ([by_id[int(override)]] if override and str(override).isdigit() and
                       int(override) in by_id else list({account.pk: account for account in
                       by_name.get(_account_name(mapped["account"]), [])}.values()))
            if len(matches) != 1:
                raise ValidationError(f"الحساب «{mapped['account']}» غير موجود أو اسمه مكرر؛ اربطه بدليل الحسابات في شاشة المطابقة أو استخدم رمزه.")
            account = matches[0]
            if job.column_map["_context"].get("layout") == "main_only" and account.level != 5:
                raise ValidationError("هذا نموذج حسابات رئيسية فقط؛ اختر حسابًا من المستوى الخامس.")
            if account.pk in seen:
                raise ValidationError("الحساب مكرر داخل الملف.")
            seen.add(account.pk)
            if account.level == 6 and (not account.parent or account.parent.level != 5 or not account.parent.is_active):
                raise ValidationError("الحساب التحليلي يجب أن يتبع حساب مصروف نشطًا من المستوى الخامس.")
            months = [_amount(mapped[field], label) for field, label in FIELDS[1:13]]
            total = _amount(mapped["total"], "الإجمالي")
            if sum(months) != total:
                raise ValidationError("الإجمالي لا يساوي مجموع الشهور الاثني عشر.")
            if not is_budget and any(value for index, value in enumerate(months, 1) if index != scope.month):
                raise ValidationError(f"ملف المصروفات لهذه الفترة يسمح بمبلغ في الشهر {scope.month} فقط.")
            row["ready"] = {"account_id": account.pk, "main_id": account.parent_id if account.level == 6 else account.pk,
                            "level": account.level, "months": [str(value) for value in months]}
            row["valid"] = True
        except ValidationError as exc:
            row["errors"].extend(exc.messages)

    parents = {row["ready"].get("account_id"): row for row in job.rows
               if row["valid"] and row["ready"].get("level") == 5}
    for parent_id, parent in parents.items():
        children = [row for row in job.rows if row["valid"] and row["ready"].get("level") == 6
                    and row["ready"]["main_id"] == parent_id]
        if not children:
            continue
        for index, month in enumerate(MONTH_FIELDS):
            if Decimal(parent["ready"]["months"][index]) != sum(Decimal(child["ready"]["months"][index]) for child in children):
                parent["errors"].append(f"{dict(FIELDS)[month]}: إجمالي الرئيسي لا يساوي مجموع التحليليات التابعة له.")
                parent["valid"] = False
        if parent["valid"]:
            parent["ready"]["skip"] = True
            parent["warnings"].append("صف إجمالي رئيسي؛ لا يُسجَّل مرة ثانية منعًا لمضاعفة المبالغ.")

    for row in job.rows:
        if not row["valid"] or row["ready"].get("skip"):
            continue
        ready = row["ready"]
        if is_budget:
            duplicate = BudgetLine.objects.filter(version=scope, department=department,
                account_id=ready["main_id"], analytical_account_id=ready["account_id"] if ready["level"] == 6 else None).exists()
            if ready["level"] == 6:
                duplicate |= BudgetLine.objects.filter(version=scope, department=department,
                    account_id=ready["main_id"], analytical_account__isnull=True).exists()
            else:
                duplicate |= BudgetLine.objects.filter(version=scope, department=department,
                    account_id=ready["main_id"], analytical_account__isnull=False).exists()
        else:
            duplicate = Expense.objects.filter(payment_reference=_reference(scope, department,
                ready["account_id"])).exists()
            if ready["level"] == 6:
                duplicate |= Expense.objects.filter(payment_reference=_reference(scope, department, ready["main_id"])).exists()
            else:
                duplicate |= Expense.objects.filter(
                    payment_reference__startswith=f"MATRIX-{scope.pk}-{department.pk}-",
                    analytical_account__parent_id=ready["main_id"]).exists()
        if duplicate:
            row["duplicate"] = True
            row["valid"] = False
            row["errors"].append("سبق تسجيل هذا الحساب لهذه الفترة؛ لا يُستبدل تلقائيًا.")

    errors = [{"row": row["index"], "message": error} for row in job.rows for error in row["errors"]]
    warnings = [{"row": row["index"], "message": warning} for row in job.rows for warning in row["warnings"]]
    job.summary = {"total": len(job.rows), "valid": sum(row["valid"] for row in job.rows),
                   "invalid": sum(bool(row["errors"]) and not row["duplicate"] for row in job.rows),
                   "duplicate": sum(row["duplicate"] for row in job.rows), "errors": len(errors),
                   "warnings": len(warnings), "error_list": errors[:300], "warning_list": warnings[:300]}
    job.status = ImportJob.STATUS_VALIDATED
    job.log("validation_completed", **{key: job.summary[key] for key in ("total", "valid", "invalid", "duplicate")})
    job.save()
    return job.summary


@transaction.atomic
def import_matrix(job, user, request=None):
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status != ImportJob.STATUS_VALIDATED or not job.summary.get("valid") or job.summary.get("errors"):
        raise ValidationError("صحّح جميع صفوف الملف وأعد التحقق قبل اعتماد الاستيراد.")
    department, scope = _context(job)
    is_budget = job.target == ImportJob.TARGET_BUDGET_MATRIX
    if is_budget:
        scope = BudgetVersion.objects.select_for_update().get(pk=scope.pk)
        if scope.status != BudgetVersion.STATUS_DRAFT:
            raise ValidationError("نسخة الموازنة أصبحت معتمدة.")
    else:
        scope = MonthlyPeriod.objects.select_for_update().get(pk=scope.pk)
        require_period_postable(user, scope, purpose="استيراد مصروفات مجمعة", request=request)
    summary = validate_matrix(job, user, request=request)
    if summary["errors"] or not summary["valid"]:
        raise ValidationError("تغيّرت البيانات بعد التحقق؛ أعد مراجعة الملف ثم أكد الاستيراد.")
    job.refresh_from_db()
    if not is_budget:
        if not any(row["valid"] and not row["ready"].get("skip") and
                   Decimal(row["ready"]["months"][scope.month - 1]) for row in job.rows):
            raise ValidationError("لا يحتوي الشهر المختار على مصروفات فعلية موجبة للاستيراد.")
    imported = 0
    for row in job.rows:
        if not row["valid"] or row["ready"].get("skip"):
            continue
        ready = row["ready"]
        account = Account.objects.get(pk=ready["account_id"])
        main = account.parent if ready["level"] == 6 else account
        if (not account.is_active or account.account_type != "expense" or
                account.level != ready["level"] or not main.is_active or
                main.level != 5 or main.pk != ready["main_id"]):
            raise ValidationError("تغيّرت حالة أحد الحسابات بعد التحقق؛ أعد التحقق من الملف.")
        if is_budget:
            if BudgetLine.objects.filter(version=scope, department=department, account=main,
                    analytical_account=account if ready["level"] == 6 else None).exists():
                raise ValidationError("سطر موازنة مسجل مسبقًا لهذا الحساب؛ أعد التحقق.")
            line = BudgetLine(version=scope, department=department, account=main,
                              analytical_account=account if ready["level"] == 6 else None,
                              created_by=user, updated_by=user,
                              **dict(zip(MONTH_FIELDS, (Decimal(value) for value in ready["months"]))))
            line.full_clean()
            line.save()
            imported += 1
        else:
            amount = Decimal(ready["months"][scope.month - 1])
            if not amount:
                continue
            reference = _reference(scope, department, account.pk)
            if Expense.objects.filter(payment_reference=reference).exists():
                raise ValidationError("مصروف مجمع مسجل مسبقًا لهذا الحساب والفترة؛ أعد التحقق.")
            expense = Expense(period=scope, expense_date=scope.end_date, department=department,
                              account=main, analytical_account=account if ready["level"] == 6 else None,
                              amount=amount, description=f"مصروفات شهرية مجمعة — {scope}",
                              payment_reference=reference,
                              expense_number=next_expense_number(scope.end_date), created_by=user, updated_by=user)
            expense.full_clean()
            expense.save()
            imported += 1
    job.status = ImportJob.STATUS_COMPLETED
    job.imported_count = imported
    job.log("import_completed", imported=imported, skipped=len(job.rows) - imported)
    job.save()
    log_action(action="data_imported", entity_type="import_job", entity_id=job.pk,
               diff={"target": job.target, "file": job.original_filename, "imported": imported},
               request=request, actor=user)
    return imported
